# -*- coding: utf-8 -*-
"""Ürün ağaçlarından elektrik listesi, kritik parça Excel/HTML üretimi."""
from __future__ import annotations

import json
import os
import re
from collections import defaultdict
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
RESEARCH = os.path.dirname(BASE)


def find_folder() -> str:
    for name in os.listdir(RESEARCH):
        p = os.path.join(RESEARCH, name)
        if os.path.isdir(p) and "PAR" in name.upper() and "KR" in name.upper():
            return p
    return BASE


FOLDER = find_folder()


def norm_code(raw) -> str | None:
    if raw is None:
        return None
    s = str(raw).strip().upper()
    if not s or s == "NONE":
        return None
    compact = re.sub(r"\s+", "", s)
    m = re.match(r"^(\d{2})(\d{3,6})$", compact)
    if m:
        return f"{m.group(1)} {m.group(2).zfill(5)}"
    m = re.match(r"^(\d{2})\s+(\d+)$", s)
    if m:
        return f"{m.group(1)} {m.group(2).zfill(5)}"
    return s


ELEC_KEYWORDS = re.compile(
    r"POMPA|SWITCH|TERMOKUPL|TERMİK|TERMIK|KONTAKT|RÖLE|ROLE|INVERT|VFD|OMRON|DELTA|"
    r"FAZ KORUMA|EMNİYET|EMNIYET|BUTON|KABLO|TTR|SİLİKON|SILIKON|FİŞ|FIS|PRİZ|PRIZ|"
    r"PLC|HMI|SENSÖR|SENSOR|SEVİYE|SEVIYE|REZİSTANS|REZISTANS|REDÜKTÖR MOTOR|"
    r"REDUKTOR MOTOR|MOTORLU|İNVERT|INVERT|ELEKTRİK|ELEKTRIK|PANO|TMŞ|LIMIT|INTERLOCK|"
    r"LAMBA|SİNYAL|SINYAL|ZAMAN RÖLE|ZAMAN ROLE|MKR|G9S|F3S|WATER SOUND|ATLANT|GOULDS|"
    r"380V|220V|50HZ|AMORTİSÖR|AMORTISOR|TORBA FİLTRE|TORBA FILTRE",
    re.I,
)

PROCESS_KEYWORDS = re.compile(
    r"EMİŞ FİLTRE|EMIS FILTRE|ÖN FİLTRE|ON FILTRE|ÖN FİLTRE NS|ON FILTRE NS|"
    r"REZİSTANS KOMP|REZISTANS KOMP|SEVİYE SENS|SEVIYE SENS|SEPET DURDURMA|"
    r"YAĞ SIYIRICI|YAG SIYIRICI|SALYANGOZLU FAN|KURUTMA KOMP|POMPA KOMP|"
    r"BORULAMA KOMP|KAPAK KOMP|TANK KOMP|SEPET KOMP|NOZZLE",
    re.I,
)

CRITICAL_ELEC = re.compile(
    r"POMPA|SWITCH|TERMOKUPL|TERMİK|TERMIK|KONTAKT|RÖLE|ROLE|INVERT|VFD|"
    r"FAZ KORUMA|EMNİYET|EMNIYET|ACİL|ACIL|INTERLOCK|LIMIT|SEVİYE BEK|SEVIYE BEK|"
    r"SEVİYE SENS|SEVIYE SENS|PASLANMAZ SEVİYE|REDÜKTÖR MOTOR|REDUKTOR MOTOR|"
    r"REDÜKTÖR EN:|REDUKTOR EN:|ELEKTRİK KUTUSU|ELEKTRIK KUTUSU|TMŞ|MKR|G9S|F3S|"
    r"TORBA FİLTRE|TORBA FILTRE",
    re.I,
)

CRITICAL_MECH = re.compile(
    r"EMİŞ FİLTRE KOMP|EMIS FILTRE KOMP|ÖN FİLTRE.*KOMP|ON FILTRE.*KOMP|"
    r"REZİSTANS KOMP|REZISTANS KOMP|SEPET DURDURMA.*SENS|YAĞ SIYIRICI|YAG SIYIRICI|"
    r"SALYANGOZLU FAN|KURUTMA KOMP|SEVİYE SENS|SEVIYE SENS",
    re.I,
)

EXCLUDE_CRITICAL = re.compile(
    r"CİVATA|CIVATA|SOMUN|PUL |KABLO |KABLO RAKORU|OSPA-|SPİRAL|SPIRAL|ETİKET|ETIKET|"
    r"KELEPÇE|KELEPCE|HORTUM|MAPa|MAPA|O-RING|RULMAN|KEÇE|KECE|DİRSEK|DIRSEK|PVC|"
    r"SEPET KOMP|ARABA\+SEPET|GÖVDE KOMP|GOVDE KOMP|TANK KOMP|KAPAK KOMP|"
    r"ÇASİ KOMP|CASI KOMP| BORULAMA KOMP|POMPA BAĞLANTI| NS S$| \d+B S$",
    re.I,
)


def is_electrical(code: str | None, name: str) -> bool:
    if code and code.startswith("10"):
        return True
    return bool(ELEC_KEYWORDS.search(name) or PROCESS_KEYWORDS.search(name))


WEAR_DEPOT = re.compile(
    r"ÖN FİLTRE|ON FILTRE|EMİŞ FİLTRE|EMIS FILTRE|NOZZLE|TORBA FİLTRE|TORBA FILTRE|"
    r"FİLTRE KOMP|FILTRE KOMP|KEÇE|KECE|AMORTİSÖR|AMORTISOR",
    re.I,
)

OPERATION_STOP = re.compile(
    r"POMPA|SWITCH|INTERLOCK|TERMOKUPL|SEVİYE BEK|SEVIYE BEK|SEVİYE SENS|SEVIYE SENS|"
    r"REDÜKTÖR MOTOR|REDUKTOR MOTOR|REDÜKTÖR EN|REDUKTOR EN|INVERT|VFD|"
    r"ELEKTRİK KUTUSU|ELEKTRIK KUTUSU|FAZ KORUMA|EMNİYET|EMNIYET|ACİL|ACIL|"
    r"KURUTMA KOMP|YAĞ SIYIRICI|YAG SIYIRICI|SALYANGOZLU FAN|REZİSTANS KOMP|REZISTANS KOMP|"
    r"SEPET DURDURMA|KONTAKT|TERMİK|TERMIK",
    re.I,
)


def criticality_axes(name: str, reason: str | None) -> tuple[str, str, str]:
    """İşletme (duruş) vs depo (aşınan/tüketim) — bağımsız iki eksen."""
    isletme = bool(reason and OPERATION_STOP.search(name))
    if reason and re.search(r"Güvenlik|Proses duruşu|Pano", reason):
        isletme = True
    if reason and re.search(r"Mekanik proses", reason):
        if re.search(r"FİLTRE|FILTRE", name, re.I):
            isletme = True
        elif re.search(r"KURUTMA|YAĞ SIYIRICI|REZİSTANS|SEPET DURDURMA|SEVİYE", name, re.I):
            isletme = True

    depo = bool(WEAR_DEPOT.search(name))
    if re.search(r"SWITCH|TERMOKUPL|SEVİYE BEK|SEVIYE BEK|FAZ KORUMA|ÖN FİLTRE|EMİŞ FİLTRE", name, re.I):
        depo = True

    ie = "Evet" if isletme else "Hayır"
    de = "Evet" if depo else "Hayır"
    if isletme and depo:
        ozet = "Duruş + depo stoku"
    elif isletme:
        ozet = "Yalnızca işletme (duruş)"
    elif depo:
        ozet = "Yalnızca depo (aşınan/tüketim)"
    else:
        ozet = "Tanımsız"
    return ie, de, ozet


def suggested_stock(isletme: str, depo: str, name: str) -> int:
    if depo == "Evet":
        return 1
    if isletme == "Evet" and re.search(
        r"SWITCH|TERMOKUPL|SEVİYE BEK|SEVIYE BEK|FAZ KORUMA|TERMİK|TERMIK", name, re.I
    ):
        return 1
    return 0


def critical_reason(code: str | None, name: str) -> str | None:
    if EXCLUDE_CRITICAL.search(name):
        return None
    if CRITICAL_ELEC.search(name):
        if re.search(r"SWITCH|INTERLOCK|EMNİYET|EMNIYET|ACİL|ACIL|FAZ KORUMA|G9S|F3S", name, re.I):
            return "Güvenlik / emniyet zinciri"
        if re.search(r"POMPA|REDÜKTÖR MOTOR|REDUKTOR MOTOR|REDÜKTÖR EN|REDUKTOR EN|INVERT|VFD", name, re.I):
            return "Proses duruşu — tahrik / pompa"
        if re.search(r"TERMOKUPL|REZİSTANS|REZISTANS|SEVİYE|SEVIYE|TORBA FİLTRE", name, re.I):
            return "Proses duruşu — ısıtma / seviye / filtrasyon"
        if re.search(r"ELEKTRİK KUTUSU|ELEKTRIK KUTUSU|KONTAKT|TERMİK|TERMIK", name, re.I):
            return "Pano / güç ve koruma"
        return "Elektrik / otomasyon — kritik"
    if code and code.startswith("07") and CRITICAL_MECH.search(name):
        return "Mekanik proses — filtre / ısıtma / sepet / kurutma"
    if code and code.startswith("10") and re.search(r"POMPA|SWITCH|TERMOKUPL|KONTAKT|REDÜKTÖR MOTOR", name, re.I):
        return critical_reason(code, name) or "Elektrik stok — kritik"
    return None


def parse_variant_columns(ws, header_row: int = 3) -> dict[str, list[dict]]:
    row1 = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0]
    variants: list[tuple[int, str]] = []
    for i, v in enumerate(row1):
        if v and isinstance(v, str) and len(v.strip()) > 2:
            variants.append((i, v.strip()))
    data: dict[str, list[dict]] = {v[1]: [] for v in variants}
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        for col_idx, vname in variants:
            if col_idx >= len(row):
                continue
            code = norm_code(row[col_idx])
            name = row[col_idx + 1] if col_idx + 1 < len(row) else None
            qty = row[col_idx + 2] if col_idx + 2 < len(row) else None
            if code and name:
                data[vname].append(
                    {"code": code, "name": str(name).strip(), "qty": qty, "variant": vname}
                )
    return data


def parse_kbn_wide(path: str) -> dict[str, list[dict]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    headers = rows[0]
    blocks: list[tuple[int, str]] = []
    for i, h in enumerate(headers):
        if h and isinstance(h, str) and "KBN" in h.upper():
            blocks.append((i, h.strip()))
    data: dict[str, list[dict]] = {b[1]: [] for b in blocks}
    for row in rows[3:]:
        for col, vname in blocks:
            if col + 4 >= len(row):
                continue
            code = norm_code(row[col + 2])
            name = row[col + 3]
            qty = row[col + 4]
            if code and name:
                data[vname].append(
                    {
                        "code": code,
                        "name": str(name).strip(),
                        "qty": qty,
                        "variant": vname,
                    }
                )
    return data


def load_groups() -> dict[str, dict[str, list[dict]]]:
    files = os.listdir(FOLDER)
    machine = next(f for f in files if f.endswith(".xlsx") and "MAK" in f.upper())
    kbn1 = next(f for f in files if "1B" in f and f.endswith(".xlsx"))
    kbn2 = next(f for f in files if "2B" in f and f.endswith(".xlsx"))
    wb = openpyxl.load_workbook(os.path.join(FOLDER, machine), read_only=True, data_only=True)
    groups = {
        "LYM": parse_variant_columns(wb["LYM"]),
        "VDL": parse_variant_columns(wb["VDL"]),
    }
    wb.close()
    groups["KBN"] = {**parse_kbn_wide(os.path.join(FOLDER, kbn1)), **parse_kbn_wide(os.path.join(FOLDER, kbn2))}
    return groups


def aggregate_unique(items_by_variant: dict[str, list[dict]]) -> dict[str, dict]:
    agg: dict[str, dict] = {}
    for variant, items in items_by_variant.items():
        for it in items:
            c = it["code"]
            if c not in agg:
                agg[c] = {
                    "code": c,
                    "name": it["name"],
                    "variants": set(),
                    "max_qty": it["qty"],
                }
            agg[c]["variants"].add(variant)
            try:
                q = float(it["qty"]) if it["qty"] is not None else 0
                mq = float(agg[c]["max_qty"]) if agg[c]["max_qty"] is not None else 0
                if q > mq:
                    agg[c]["max_qty"] = it["qty"]
            except (TypeError, ValueError):
                pass
    return agg


def load_komple_bom() -> dict[str, dict]:
    path = os.path.join(FOLDER, "komple-alt-malzemeler.json")
    if not os.path.isfile(path):
        return {}
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    raw = data.get("kompleler", data)
    out: dict[str, dict] = {}
    for key, val in raw.items():
        if key.startswith("_"):
            continue
        nk = norm_code(key) or key.strip().upper()
        out[nk] = val if isinstance(val, dict) else {"alt_malzemeler": val}
    return out


def _child_axes(name: str, parent: dict) -> tuple[str, str, str, int]:
    """Alt kalem kritikliği — aşınan alt parçalar depo; tahrik/ısı parçaları duruş."""
    isletme = "Hayır"
    depo = "Evet"
    if re.search(r"REZİSTANS 220|REZISTANS 220|8250W", name, re.I):
        isletme = "Evet"
    if re.search(
        r"POMPA|BLOWER|FAN |REDÜKTÖR|REDUKTOR|KONTAKT|TERMİK|TERMIK|INTERLOCK|SWITCH|LIMIT",
        name,
        re.I,
    ):
        isletme = "Evet"
    if re.search(r"KEÇE|KECE|TEFLON|CONTA|KLİNGR|KLINGR|SOMUN", name, re.I):
        depo = "Evet"
        isletme = "Hayır"
    if isletme == "Evet" and depo == "Evet":
        ozet = "Duruş + depo stoku"
    elif isletme == "Evet":
        ozet = "Yalnızca işletme (duruş)"
    else:
        ozet = "Yalnızca depo (aşınan/tüketim)"
    stock = suggested_stock(isletme, depo, name)
    return isletme, depo, ozet, stock


def _make_child_row(
    ck: str,
    cname: str,
    seviye: str,
    ust: str,
    root_row: dict,
    parent_komple: str,
    siparis: str,
    inherit: dict,
) -> dict:
    if seviye == "Alt komple":
        ci = inherit["İşletme kritiği (duruş)"]
        cd = inherit["Depo stoku (aşınan/tüketim)"]
        co = inherit["Kritiklik özeti"]
        cstock = inherit["Önerilen stok"]
        gerekce = inherit.get("Teknik gerekçe (alan)", "")
    else:
        ci, cd, co, cstock = _child_axes(cname, inherit)
        gerekce = f"Alt malzeme — {parent_komple} komple alternatifi"
    return {
        "Sipariş kodu": ck,
        "Parça adı": cname,
        "Seviye": seviye,
        "Üst komple kodu": ust,
        "Sipariş seçeneği": siparis,
        "İşletme kritiği (duruş)": ci,
        "Depo stoku (aşınan/tüketim)": cd,
        "Kritiklik özeti": co,
        "Önerilen stok": cstock,
        "Kullanıldığı varyantlar": root_row["Kullanıldığı varyantlar"],
        "Montaj adedi (max)": "—",
        "Teknik gerekçe (alan)": gerekce,
        "Grup": root_row.get("Grup", ""),
    }


def _append_bom_tree(
    expanded: list[dict],
    root_row: dict,
    code: str,
    entry: dict,
    bom: dict[str, dict],
    inherit: dict,
) -> None:
    for child in entry.get("alt_malzemeler", []):
        ck = norm_code(child.get("kod")) or child.get("kod", "")
        cname = child.get("ad", "")
        sub = bom.get(ck)
        if sub and sub.get("alt_malzemeler"):
            expanded.append(
                _make_child_row(
                    ck,
                    cname,
                    "Alt komple",
                    code,
                    root_row,
                    code,
                    "Komple veya alt kalem",
                    inherit,
                )
            )
            _append_bom_tree(expanded, root_row, ck, sub, bom, inherit)
        else:
            expanded.append(
                _make_child_row(
                    ck,
                    cname,
                    "Alt malzeme",
                    code,
                    root_row,
                    code,
                    f"Alt kalem ({code} komplesi)",
                    inherit,
                )
            )


def expand_komple_subparts(rows: list[dict], bom: dict[str, dict]) -> list[dict]:
    if not bom:
        return rows
    expanded: list[dict] = []
    for row in rows:
        code = norm_code(row["Sipariş kodu"]) or row["Sipariş kodu"]
        entry = bom.get(code)
        parent = {
            **row,
            "Seviye": "Komple",
            "Üst komple kodu": "",
            "Sipariş seçeneği": "Komple veya alt kalem" if entry else "",
        }
        expanded.append(parent)
        if entry:
            _append_bom_tree(expanded, row, code, entry, bom, parent)
    return expanded


def build_critical_rows(group: str, agg: dict[str, dict]) -> list[dict]:
    rows = []
    for c, info in sorted(agg.items(), key=lambda x: x[0]):
        reason = critical_reason(c, info["name"])
        if not reason:
            continue
        isletme, depo, ozet = criticality_axes(info["name"], reason)
        stock = suggested_stock(isletme, depo, info["name"])
        rows.append(
            {
                "Sipariş kodu": c,
                "Parça adı": info["name"],
                "Seviye": "Komple",
                "Üst komple kodu": "",
                "Sipariş seçeneği": "",
                "İşletme kritiği (duruş)": isletme,
                "Depo stoku (aşınan/tüketim)": depo,
                "Kritiklik özeti": ozet,
                "Önerilen stok": stock,
                "Kullanıldığı varyantlar": ", ".join(sorted(info["variants"])),
                "Montaj adedi (max)": info["max_qty"],
                "Teknik gerekçe (alan)": reason,
                "Grup": group,
            }
        )
    bom = load_komple_bom()
    return expand_komple_subparts(rows, bom)


def write_md(groups: dict, out_path: str) -> None:
    lines = [
        "# LYM · VDL · KBN — Ürün Ağacı Parça Listeleri",
        "",
        f"Kaynak: `MAKİNE ÜRÜN AĞAÇLARI.xlsx`, `KBN 1B/2B GRUBU ÜRÜN AĞAÇLARI.xlsx` · Oluşturma: {date.today().isoformat()}",
        "",
        "Bu dosya ürün ağaçlarından türetilmiş **benzersiz stok kodu** listesidir. "
        "**Elektrik / otomasyon filtresi** (stok kodu `10 …` veya elektrik/proses anahtar kelimesi) ayrı bölümde verilmiştir.",
        "",
    ]
    for gname in ("LYM", "VDL", "KBN"):
        agg = aggregate_unique(groups[gname])
        elec = {c: v for c, v in agg.items() if is_electrical(c, v["name"])}
        lines += [
            f"## {gname} — Tüm parçalar ({len(agg)} benzersiz kod)",
            "",
            "| Sipariş kodu | Parça adı | Varyant sayısı |",
            "|---|---|---|",
        ]
        for c, info in sorted(agg.items(), key=lambda x: x[0]):
            lines.append(f"| {c} | {info['name']} | {len(info['variants'])} |")
        lines += [
            "",
            f"### {gname} — Elektrik / otomasyon filtresi ({len(elec)} kod)",
            "",
            "| Sipariş kodu | Parça adı | Varyantlar |",
            "|---|---|---|",
        ]
        for c, info in sorted(elec.items(), key=lambda x: x[0]):
            vars_short = ", ".join(sorted(info["variants"]))
            if len(vars_short) > 80:
                vars_short = vars_short[:77] + "…"
            lines.append(f"| {c} | {info['name']} | {vars_short} |")
        lines.append("")

    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))


def write_excel(all_critical: dict[str, list[dict]], out_path: str) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)
    cols = [
        "Seviye",
        "Üst komple kodu",
        "Sipariş kodu",
        "Parça adı",
        "Sipariş seçeneği",
        "İşletme kritiği (duruş)",
        "Depo stoku (aşınan/tüketim)",
        "Kritiklik özeti",
        "Önerilen stok",
        "Montaj adedi (max)",
        "Kullanıldığı varyantlar",
        "Teknik gerekçe (alan)",
    ]
    child_font = Font(color="475569")
    komple_fill = PatternFill("solid", fgColor="F0F9FF")
    alt_komple_fill = PatternFill("solid", fgColor="ECFDF5")
    for sheet_name in ("LYM", "VDL", "KBN"):
        ws = wb.create_sheet(sheet_name)
        ws.append(cols)
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in all_critical[sheet_name]:
            ws.append([row.get(c, "") for c in cols])
            r = ws.max_row
            if row.get("Seviye") == "Alt malzeme":
                for cell in ws[r]:
                    cell.font = child_font
                ws.cell(r, 4).alignment = Alignment(indent=3, wrap_text=True)
            elif row.get("Seviye") == "Alt komple":
                for cell in ws[r]:
                    cell.fill = alt_komple_fill
                ws.cell(r, 4).alignment = Alignment(indent=1, wrap_text=True)
            elif row.get("Sipariş seçeneği"):
                for cell in ws[r]:
                    cell.fill = komple_fill
        for i, w in enumerate((12, 14, 14, 46, 28, 10, 12, 22, 10, 12, 38, 30), start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
    wb.save(out_path)


def write_html(all_critical: dict[str, list[dict]], out_path: str) -> None:
    data_json = json.dumps(all_critical, ensure_ascii=False)
    html = f"""<!DOCTYPE html>
<html lang="tr">
<head>
  <meta charset="utf-8"/>
  <meta name="viewport" content="width=device-width, initial-scale=1"/>
  <title>LYM · VDL · KBN Kritik Parça Listeleri</title>
  <style>
    :root {{
      --bg:#eef1f6; --surface:#fff; --border:#d4dbe6; --text:#0f172a; --muted:#64748b;
      --accent:#2563eb; --accent-soft:#dbeafe;
      --guvenlik:#dc2626; --guvenlik-bg:#fef2f2;
      --tahrik:#0369a1; --tahrik-bg:#e0f2fe;
      --olcum:#7c3aed; --olcum-bg:#f3e8ff;
      --pano:#4338ca; --pano-bg:#eef2ff;
      --mekanik:#0d9488; --mekanik-bg:#ccfbf1;
      --elektrik:#475569; --elektrik-bg:#f1f5f9;
      --radius:12px;
      --shadow:0 4px 24px rgba(15,23,42,.06);
    }}
    * {{ box-sizing:border-box; }}
    body {{ margin:0; font-family:"Segoe UI",system-ui,sans-serif; background:var(--bg); color:var(--text); line-height:1.45; }}
    header {{
      background:linear-gradient(120deg,#0f172a 0%,#1e3a8a 55%,#2563eb 100%);
      color:#fff; padding:1.5rem 1.75rem 1.35rem;
      box-shadow:0 2px 12px rgba(0,0,0,.15);
    }}
    header h1 {{ margin:0 0 .4rem; font-size:1.45rem; font-weight:650; letter-spacing:-.02em; }}
    header p {{ margin:0; opacity:.88; font-size:.9rem; max-width:52rem; }}
    .topbar {{
      position:sticky; top:0; z-index:20;
      background:rgba(238,241,246,.92); backdrop-filter:blur(10px);
      border-bottom:1px solid var(--border);
      padding:.65rem 1.25rem;
      display:flex; flex-wrap:wrap; gap:.5rem; align-items:center;
    }}
    .topbar .tabs {{ display:flex; gap:.4rem; flex-wrap:wrap; }}
    .tab {{
      border:1px solid var(--border); background:var(--surface);
      padding:.45rem .95rem; border-radius:999px; cursor:pointer; font-weight:600; font-size:.85rem;
      transition:all .15s ease;
    }}
    .tab:hover {{ border-color:#94a3b8; }}
    .tab.active {{ background:var(--accent); color:#fff; border-color:var(--accent); box-shadow:0 2px 8px rgba(37,99,235,.35); }}
    .tab .n {{ opacity:.75; font-weight:500; margin-left:.35rem; }}
    main {{ max-width:1480px; margin:0 auto; padding:1rem 1.25rem 2.5rem; }}
    .stats {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(140px,1fr)); gap:.65rem; margin-bottom:1rem; }}
    .stat {{
      background:var(--surface); border:1px solid var(--border); border-radius:var(--radius);
      padding:.65rem .85rem; box-shadow:var(--shadow);
    }}
    .stat .label {{ font-size:.72rem; text-transform:uppercase; letter-spacing:.04em; color:var(--muted); font-weight:600; }}
    .stat .val {{ font-size:1.35rem; font-weight:700; margin-top:.15rem; }}
    .panel {{
      background:var(--surface); border:1px solid var(--border); border-radius:var(--radius);
      box-shadow:var(--shadow); overflow:hidden;
    }}
    .toolbar {{
      padding:.85rem 1rem; border-bottom:1px solid var(--border);
      display:flex; flex-wrap:wrap; gap:.65rem; align-items:center;
    }}
    .search-wrap {{ flex:1; min-width:200px; position:relative; }}
    .search-wrap svg {{ position:absolute; left:.65rem; top:50%; transform:translateY(-50%); opacity:.45; }}
    #q {{
      width:100%; padding:.55rem .75rem .55rem 2.1rem;
      border:1px solid var(--border); border-radius:10px; font-size:.9rem;
      background:#fafbfc;
    }}
    #q:focus {{ outline:2px solid var(--accent-soft); border-color:var(--accent); background:#fff; }}
    .filters {{ display:flex; flex-wrap:wrap; gap:.35rem; align-items:center; }}
    .filters span {{ font-size:.75rem; color:var(--muted); font-weight:600; margin-right:.25rem; }}
    .chip {{
      border:1px solid var(--border); background:#f8fafc; border-radius:999px;
      padding:.25rem .65rem; font-size:.75rem; font-weight:600; cursor:pointer; user-select:none;
    }}
    .chip.on {{ background:var(--accent-soft); border-color:#93c5fd; color:#1d4ed8; }}
    .meta {{ font-size:.82rem; color:var(--muted); white-space:nowrap; }}
    .table-wrap {{ overflow:auto; max-height:calc(100vh - 320px); }}
    table {{ width:100%; border-collapse:separate; border-spacing:0; font-size:.8125rem; table-layout:fixed; }}
    col.col-tree {{ width:168px; }}
    col.col-code {{ width:118px; }}
    col.col-name {{ width:min(280px, 22vw); }}
    thead th {{
      position:sticky; top:0; z-index:5;
      background:#f8fafc; border-bottom:2px solid var(--border);
      padding:.65rem .75rem; text-align:left; font-size:.72rem; text-transform:uppercase;
      letter-spacing:.035em; color:#475569; font-weight:700;
      box-shadow:0 1px 0 var(--border);
    }}
    tbody td {{ padding:.6rem .75rem; border-bottom:1px solid #eef2f7; vertical-align:top; }}
    tbody tr {{ transition:background .12s; }}
    tbody tr:hover td {{ background:#f8fafc; }}
    tbody tr.row-alt td {{ background:#fbfcfe; }}
    tbody tr[data-alan="guvenlik"] {{ border-left:3px solid var(--guvenlik); }}
    tbody tr[data-alan="tahrik"] {{ border-left:3px solid var(--tahrik); }}
    tbody tr[data-alan="olcum"] {{ border-left:3px solid var(--olcum); }}
    tbody tr[data-alan="pano"] {{ border-left:3px solid var(--pano); }}
    tbody tr[data-alan="mekanik"] {{ border-left:3px solid var(--mekanik); }}
    tbody tr[data-alan="elektrik"] {{ border-left:3px solid var(--elektrik); }}
    tbody tr[data-alan="diger"] {{ border-left:3px solid #cbd5e1; }}
    tbody tr[data-tip="durus"] {{ box-shadow:inset 3px 0 0 #dc2626; }}
    tbody tr[data-tip="depo"] {{ box-shadow:inset 3px 0 0 #d97706; }}
    tbody tr[data-tip="both"] {{ box-shadow:inset 3px 0 0 #7c3aed; }}
    .code {{
      font-family:Consolas,"Cascadia Mono",monospace; font-size:.78rem; font-weight:600;
      background:#f1f5f9; border:1px solid #e2e8f0; padding:.2rem .45rem; border-radius:6px;
      white-space:nowrap;
    }}
    .part-name {{ font-weight:500; max-width:28rem; }}
    .class-cell {{ display:flex; flex-direction:column; gap:.35rem; min-width:9rem; }}
    .badge-row {{ display:flex; flex-wrap:wrap; gap:.3rem; }}
    .badge {{
      display:inline-flex; align-items:center; gap:.25rem;
      padding:.18rem .5rem; border-radius:6px; font-size:.68rem; font-weight:700;
      letter-spacing:.02em; border:1px solid transparent;
    }}
    .badge-durus {{ background:#fee2e2; color:#991b1b; border-color:#fecaca; }}
    .badge-depo {{ background:#fef3c7; color:#92400e; border-color:#fde68a; }}
    .badge-both {{ background:#ede9fe; color:#5b21b6; border-color:#ddd6fe; }}
    .badge-hayir {{ background:#f1f5f9; color:#64748b; border-color:#e2e8f0; font-weight:600; }}
    .badge-alan-guvenlik {{ background:var(--guvenlik-bg); color:var(--guvenlik); border-color:#fecaca; }}
    .badge-alan-tahrik {{ background:var(--tahrik-bg); color:var(--tahrik); border-color:#bae6fd; }}
    .badge-alan-olcum {{ background:var(--olcum-bg); color:var(--olcum); border-color:#ddd6fe; }}
    .badge-alan-pano {{ background:var(--pano-bg); color:var(--pano); border-color:#c7d2fe; }}
    .badge-alan-mekanik {{ background:var(--mekanik-bg); color:var(--mekanik); border-color:#99f6e4; }}
    .badge-alan-elektrik {{ background:var(--elektrik-bg); color:var(--elektrik); border-color:#e2e8f0; }}
    .badge-alan-diger {{ background:#f1f5f9; color:#64748b; border-color:#e2e8f0; }}
    .stock {{ display:inline-block; min-width:1.75rem; text-align:center; font-weight:700; border-radius:6px; padding:.15rem .4rem; font-size:.78rem; }}
    .stock-yes {{ background:#dcfce7; color:#166534; }}
    .stock-no {{ background:#f1f5f9; color:#64748b; }}
    .qty {{ font-variant-numeric:tabular-nums; font-weight:600; }}
    .reason {{ color:#475569; font-size:.78rem; max-width:16rem; }}
    .variants {{ display:flex; flex-wrap:wrap; gap:.25rem; max-width:22rem; }}
    .vtag {{
      font-size:.68rem; background:#f1f5f9; border:1px solid #e2e8f0; color:#334155;
      padding:.12rem .4rem; border-radius:4px; line-height:1.3;
    }}
    .empty {{ padding:2.5rem; text-align:center; color:var(--muted); }}
    tr.komple-row td {{ background:#f0f9ff; }}
    tr.child-row td {{ background:#fafbfc; }}
    tr.child-row .code {{ background:#fff; border-style:dashed; }}
    .tree {{ color:#64748b; margin-right:.25rem; font-weight:700; }}
    .sip-not {{ display:block; font-size:.72rem; color:#0369a1; margin-top:.3rem; font-weight:500; }}
    .seviye-tag {{ font-size:.65rem; font-weight:700; text-transform:uppercase; letter-spacing:.04em; color:#64748b; }}
    .seviye-komple {{ color:#1d4ed8; }}
    .seviye-alt {{ color:#0d9488; }}
    .seviye-alt-komple {{ color:#047857; }}
    tr.alt-komple-row td {{ background:#ecfdf5; }}
    tr.is-hidden {{ display:none; }}
    td.tree-cell {{
      background:#f8fafc; border-right:1px solid var(--border);
      vertical-align:middle; padding:.45rem .5rem;
    }}
    .tree-inner {{ display:flex; flex-direction:column; align-items:flex-start; gap:.2rem; min-height:1.75rem; }}
    .tree-line {{ display:flex; align-items:center; gap:.25rem; width:100%; }}
    .tog {{
      flex-shrink:0; width:1.35rem; height:1.35rem; border:1px solid var(--border); border-radius:6px;
      background:#fff; cursor:pointer; font-size:.65rem; line-height:1; color:#475569;
      display:inline-flex; align-items:center; justify-content:center;
    }}
    .tog:hover {{ border-color:#94a3b8; background:#f1f5f9; }}
    .tog.expanded {{ background:var(--accent-soft); border-color:#93c5fd; color:#1d4ed8; }}
    .tog-placeholder {{ width:1.35rem; flex-shrink:0; }}
    .ust-ref {{
      font-size:.65rem; color:var(--muted); font-weight:600; line-height:1.25;
      padding-left:1.6rem; max-width:100%; word-break:break-word;
    }}
    .ust-ref code {{ font-size:.62rem; background:#e2e8f0; padding:0 .25rem; border-radius:3px; }}
    .child-count {{ font-size:.62rem; color:#64748b; margin-left:.15rem; }}
    .toolbar .chip-btn {{
      border:1px solid var(--border); background:#fff; border-radius:8px;
      padding:.35rem .65rem; font-size:.75rem; font-weight:600; cursor:pointer;
    }}
    .toolbar .chip-btn:hover {{ background:#f1f5f9; }}
    .legend {{
      display:flex; flex-wrap:wrap; gap:.5rem 1rem; padding:.65rem 1rem;
      background:#fafbfc; border-bottom:1px solid var(--border); font-size:.72rem; color:var(--muted);
    }}
    .legend i {{ display:inline-block; width:10px; height:10px; border-radius:2px; margin-right:.35rem; vertical-align:middle; }}
    @media (max-width:900px) {{
      .table-wrap {{ max-height:none; }}
      thead th {{ position:static; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>Kritik parça listeleri — LYM · VDL · KBN</h1>
    <p>İki eksen: <strong>işletme kritiği (duruş)</strong> · <strong>depo stoku (aşınan/tüketim)</strong> — IEC/IEEE 82079-1 yedek/sarf dokümantasyonu</p>
  </header>
  <div class="topbar">
    <div class="tabs" id="tabs"></div>
  </div>
  <main>
    <div class="stats" id="stats"></div>
    <div class="panel">
      <div class="legend">
        <span><i style="background:#dc2626"></i>İşletme — arızada makine çalışmaz</span>
        <span><i style="background:#d97706"></i>Depo — aşınan, sürekli stok</span>
        <span><i style="background:#7c3aed"></i>Her ikisi</span>
        <span><i style="background:#bae6fd"></i>Komple — ▶ ile alt kırılım (varsayılan kapalı)</span>
        <span style="margin-left:auto;opacity:.85">Sol sütun: ağaç / üst komple</span>
      </div>
      <div class="toolbar">
        <div class="search-wrap">
          <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><circle cx="11" cy="11" r="7"/><path d="M20 20l-3-3"/></svg>
          <input type="search" id="q" placeholder="Sipariş kodu, parça adı veya gerekçe…"/>
        </div>
        <button type="button" class="chip-btn" id="expandAll">Altları aç</button>
        <button type="button" class="chip-btn" id="collapseAll">Altları kapat</button>
        <div class="filters" id="filters"></div>
        <span class="meta" id="count"></span>
      </div>
      <div class="table-wrap">
        <table>
          <colgroup>
            <col class="col-tree"/><col class="col-code"/><col class="col-name"/>
          </colgroup>
          <thead>
            <tr>
              <th>Alt malzeme ağacı</th>
              <th>Sipariş kodu</th>
              <th>Parça adı</th>
              <th>Kritiklik tipi</th>
              <th>Stok</th>
              <th>Adet</th>
              <th>Varyantlar</th>
              <th>Teknik gerekçe</th>
            </tr>
          </thead>
          <tbody id="tbody"></tbody>
        </table>
      </div>
    </div>
  </main>
  <script>
    const DATA = {data_json};

    const ALAN_META = {{
      guvenlik: {{ label: 'Güvenlik', badge: 'badge-alan-guvenlik' }},
      tahrik: {{ label: 'Tahrik / pompa', badge: 'badge-alan-tahrik' }},
      olcum: {{ label: 'Isı / seviye', badge: 'badge-alan-olcum' }},
      pano: {{ label: 'Pano', badge: 'badge-alan-pano' }},
      mekanik: {{ label: 'Mekanik proses', badge: 'badge-alan-mekanik' }},
      elektrik: {{ label: 'Otomasyon', badge: 'badge-alan-elektrik' }},
      diger: {{ label: 'Diğer', badge: 'badge-alan-diger' }},
    }};

    let active = 'LYM';
    let tipFilter = null;
    let alanFilter = new Set();
    const expandedByGroup = {{ LYM: new Set(), VDL: new Set(), KBN: new Set() }};

    const tabsEl = document.getElementById('tabs');
    const tbody = document.getElementById('tbody');
    const statsEl = document.getElementById('stats');
    const filtersEl = document.getElementById('filters');
    const q = document.getElementById('q');
    const countEl = document.getElementById('count');

    function parseRow(r) {{
      const isletme = (r['İşletme kritiği (duruş)'] || '') === 'Evet';
      const depo = (r['Depo stoku (aşınan/tüketim)'] || '') === 'Evet';
      let tip = 'none';
      if (isletme && depo) tip = 'both';
      else if (isletme) tip = 'durus';
      else if (depo) tip = 'depo';
      const g = r['Teknik gerekçe (alan)'] || r['Kritiklik gerekçesi'] || '';
      let alan = 'diger';
      if (/Güvenlik|emniyet/i.test(g)) alan = 'guvenlik';
      else if (/tahrik|pompa/i.test(g)) alan = 'tahrik';
      else if (/ısıtma|seviye|filtrasyon/i.test(g)) alan = 'olcum';
      else if (/Pano/i.test(g)) alan = 'pano';
      else if (/Mekanik/i.test(g)) alan = 'mekanik';
      else if (/Elektrik|otomasyon/i.test(g)) alan = 'elektrik';
      return {{ isletme, depo, tip, alan, ozet: r['Kritiklik özeti'] || '' }};
    }}

    function shortVariant(s) {{
      return s
        .replace(/\\s*PARÇA YIKAMA MAKİNESİ\\s*/gi, '')
        .replace(/\\s+/g, ' ')
        .trim();
    }}

    function escapeHtml(s) {{
      return String(s ?? '')
        .replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
    }}

    ['LYM','VDL','KBN'].forEach(g => {{
      const b = document.createElement('button');
      b.type = 'button';
      b.className = 'tab';
      b.dataset.group = g;
      b.innerHTML = g + '<span class="n">' + (DATA[g]?.length || 0) + '</span>';
      b.onclick = () => {{
        active = g;
        tipFilter = null;
        alanFilter.clear();
        renderFilterChips();
        render();
      }};
      tabsEl.appendChild(b);
    }});

    function renderFilterChips() {{
      filtersEl.innerHTML = '';
      const tipLbl = document.createElement('span');
      tipLbl.textContent = 'Tip:';
      filtersEl.appendChild(tipLbl);
      [
        ['durus', 'Yalnız duruş'],
        ['depo', 'Yalnız depo'],
        ['both', 'Duruş + depo'],
      ].forEach(([key, label]) => {{
        const c = document.createElement('button');
        c.type = 'button';
        c.className = 'chip' + (tipFilter === key ? ' on' : '');
        c.textContent = label;
        c.onclick = () => {{
          tipFilter = tipFilter === key ? null : key;
          renderFilterChips();
          render();
        }};
        filtersEl.appendChild(c);
      }});
      const alanLbl = document.createElement('span');
      alanLbl.textContent = 'Alan:';
      alanLbl.style.marginLeft = '.5rem';
      filtersEl.appendChild(alanLbl);
      Object.entries(ALAN_META).forEach(([key, meta]) => {{
        if (key === 'diger') return;
        const c = document.createElement('button');
        c.type = 'button';
        c.className = 'chip' + (alanFilter.has(key) ? ' on' : '');
        c.textContent = meta.label;
        c.onclick = () => {{
          if (alanFilter.has(key)) alanFilter.delete(key);
          else alanFilter.add(key);
          c.classList.toggle('on');
          render();
        }};
        filtersEl.appendChild(c);
      }});
    }}

    function syncTabs() {{
      [...tabsEl.children].forEach(b => b.classList.toggle('active', b.dataset.group === active));
    }}

    function renderStats(rows) {{
      let durus = 0, depo = 0, both = 0, stoklu = 0;
      rows.forEach(r => {{
        const {{ tip, isletme, depo: d }} = parseRow(r);
        if (tip === 'both') both++;
        else if (tip === 'durus') durus++;
        else if (tip === 'depo') depo++;
        if (Number(r['Önerilen stok']) > 0) stoklu++;
      }});
      const cards = [
        ['Toplam parça', rows.length],
        ['İşletme (duruş)', durus + both],
        ['Depo (aşınan)', depo + both],
        ['Duruş + depo', both],
        ['Önerilen stok ≥ 1', stoklu],
      ];
      statsEl.innerHTML = cards.map(([label, val]) =>
        '<div class="stat"><div class="label">' + escapeHtml(label) + '</div><div class="val">' + val + '</div></div>'
      ).join('');
    }}

    function rowMatches(r, term, p) {{
      if (tipFilter && p.tip !== tipFilter) return false;
      if (alanFilter.size && !alanFilter.has(p.alan)) return false;
      if (!term) return true;
      return Object.values(r).some(v => String(v).toLowerCase().includes(term));
    }}

    function rowDepth(r) {{
      if (r['Seviye'] === 'Komple') return 0;
      if (r['Seviye'] === 'Alt komple') return 1;
      return 2;
    }}

    function hasDirectChildren(all, idx) {{
      const code = all[idx]['Sipariş kodu'];
      for (let j = idx + 1; j < all.length; j++) {{
        if (all[j]['Seviye'] === 'Komple' && !all[j]['Üst komple kodu']) break;
        if (all[j]['Üst komple kodu'] === code) return true;
      }}
      return false;
    }}

    function countDirectChildren(all, idx) {{
      const code = all[idx]['Sipariş kodu'];
      let n = 0;
      for (let j = idx + 1; j < all.length; j++) {{
        if (all[j]['Seviye'] === 'Komple' && !all[j]['Üst komple kodu']) break;
        if (all[j]['Üst komple kodu'] === code) n++;
      }}
      return n;
    }}

    function isTreeVisible(all, idx, exp) {{
      const r = all[idx];
      if (r['Seviye'] === 'Komple' && !r['Üst komple kodu']) return true;
      let ust = r['Üst komple kodu'];
      while (ust) {{
        if (!exp.has(ust)) return false;
        let parentUst = '';
        for (let j = idx - 1; j >= 0; j--) {{
          if (all[j]['Sipariş kodu'] === ust) {{
            parentUst = all[j]['Üst komple kodu'] || '';
            break;
          }}
        }}
        ust = parentUst;
      }}
      return true;
    }}

    function expandAncestorsForSearch(all, term, exp) {{
      if (!term) return exp;
      const merged = new Set(exp);
      all.forEach((r, i) => {{
        if (!rowMatches(r, term, parseRow(r))) return;
        let ust = r['Üst komple kodu'];
        while (ust) {{
          merged.add(ust);
          let next = '';
          for (let j = i - 1; j >= 0; j--) {{
            if (all[j]['Sipariş kodu'] === ust) {{
              next = all[j]['Üst komple kodu'] || '';
              break;
            }}
          }}
          ust = next;
        }}
      }});
      return merged;
    }}

    function toggleExpand(code) {{
      const set = expandedByGroup[active];
      if (set.has(code)) set.delete(code);
      else set.add(code);
      render();
    }}

    function buildTreeCell(all, idx, exp) {{
      const r = all[idx];
      const depth = rowDepth(r);
      const pad = 6 + depth * 14;
      const code = r['Sipariş kodu'];
      const kids = hasDirectChildren(all, idx);
      const open = exp.has(code);
      let line = '<div class="tree-line" style="padding-left:' + pad + 'px">';
      if (kids) {{
        line += '<button type="button" class="tog' + (open ? ' expanded' : '') + '" data-toggle="' + escapeHtml(code) + '" aria-expanded="' + open + '">' + (open ? '▼' : '▶') + '</button>';
        const cnt = countDirectChildren(all, idx);
        if (cnt) line += '<span class="child-count">+' + cnt + '</span>';
      }} else {{
        line += '<span class="tog-placeholder"></span>';
      }}
      line += '</div>';
      const ust = r['Üst komple kodu'];
      if (ust) {{
        line += '<div class="ust-ref" style="padding-left:' + (pad + 4) + 'px">Komple: <code>' + escapeHtml(ust) + '</code></div>';
      }} else if (kids && !open) {{
        line += '<div class="ust-ref" style="padding-left:' + (pad + 4) + 'px;color:#94a3b8">Alt kırılım kapalı</div>';
      }}
      return '<td class="tree-cell"><div class="tree-inner">' + line + '</div></td>';
    }}

    function render() {{
      syncTabs();
      const term = (q.value || '').toLowerCase().trim();
      const all = DATA[active] || [];
      const exp = expandAncestorsForSearch(all, term, expandedByGroup[active]);

      renderStats(all.filter(r => r['Seviye'] === 'Komple' && !r['Üst komple kodu']));

      const parts = [];
      let shown = 0;
      all.forEach((r, idx) => {{
        if (!isTreeVisible(all, idx, exp)) return;
        if (!rowMatches(r, term, parseRow(r))) return;
        shown++;

        const {{ isletme, depo, tip, alan, ozet }} = parseRow(r);
        const isChild = r['Seviye'] === 'Alt malzeme';
        const isAltKomple = r['Seviye'] === 'Alt komple';
        const meta = ALAN_META[alan] || ALAN_META.diger;
        const durusB = isletme ? '<span class="badge badge-durus">Duruş</span>' : '<span class="badge badge-hayir">Duruş değil</span>';
        const depoB = depo ? '<span class="badge badge-depo">Depo stoku</span>' : '<span class="badge badge-hayir">Depo değil</span>';
        const ozetB = ozet ? '<span class="badge ' + (tip === 'both' ? 'badge-both' : tip === 'durus' ? 'badge-durus' : 'badge-depo') + '">' + escapeHtml(ozet) + '</span>' : '';
        const alanBadge = '<span class="badge ' + meta.badge + '">' + meta.label + '</span>';
        const stok = Number(r['Önerilen stok']) || 0;
        const stokHtml = '<span class="stock ' + (stok > 0 ? 'stock-yes' : 'stock-no') + '">' + stok + '</span>';
        const vars = String(r['Kullanıldığı varyantlar'] || '').split(',').map(s => shortVariant(s)).filter(Boolean);
        const vHtml = vars.map(v => '<span class="vtag">' + escapeHtml(v) + '</span>').join('');
        const rowTip = tip === 'both' ? 'both' : tip === 'durus' ? 'durus' : tip === 'depo' ? 'depo' : 'none';
        const sip = r['Sipariş seçeneği'] ? '<span class="sip-not">' + escapeHtml(r['Sipariş seçeneği']) + '</span>' : '';
        const rowCls = isChild ? 'child-row' : isAltKomple ? 'alt-komple-row' : 'komple-row';

        parts.push(
          '<tr data-alan="' + alan + '" data-tip="' + rowTip + '" class="' + rowCls + '">' +
          buildTreeCell(all, idx, exp) +
          '<td><span class="code">' + escapeHtml(r['Sipariş kodu']) + '</span></td>' +
          '<td class="part-name">' + escapeHtml(r['Parça adı']) + sip + '</td>' +
          '<td><div class="class-cell"><div class="badge-row">' + durusB + depoB + '</div><div class="badge-row">' + ozetB + '</div></div></td>' +
          '<td>' + stokHtml + '</td>' +
          '<td class="qty">' + escapeHtml(r['Montaj adedi (max)']) + '</td>' +
          '<td><div class="variants">' + vHtml + '</div></td>' +
          '<td class="reason">' + alanBadge + ' · ' + escapeHtml(r['Teknik gerekçe (alan)'] || r['Kritiklik gerekçesi']) + '</td>' +
          '</tr>'
        );
      }});

      tbody.innerHTML = parts.length
        ? parts.join('')
        : '<tr><td colspan="8" class="empty">Eşleşen parça yok — arama veya filtreyi değiştirin.</td></tr>';

      tbody.querySelectorAll('[data-toggle]').forEach(btn => {{
        btn.onclick = (e) => {{
          e.stopPropagation();
          toggleExpand(btn.getAttribute('data-toggle'));
        }};
      }});

      const roots = all.filter(r => r['Seviye'] === 'Komple' && !r['Üst komple kodu']).length;
      countEl.textContent = active + ': ' + shown + ' görünür / ' + all.length + ' toplam · ' + roots + ' ana komple';
    }}

    document.getElementById('expandAll').onclick = () => {{
      const all = DATA[active] || [];
      const set = expandedByGroup[active];
      all.forEach((r, i) => {{
        if (hasDirectChildren(all, i)) set.add(r['Sipariş kodu']);
      }});
      render();
    }};
    document.getElementById('collapseAll').onclick = () => {{
      expandedByGroup[active].clear();
      render();
    }};

    renderFilterChips();
    q.addEventListener('input', render);
    render();
  </script>
</body>
</html>
"""
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(html)


def main() -> None:
    groups = load_groups()
    all_critical = {g: build_critical_rows(g, aggregate_unique(groups[g])) for g in groups}

    write_md(groups, os.path.join(FOLDER, "urun-agaci-parca-listesi.md"))
    write_excel(all_critical, os.path.join(FOLDER, "Kritik-Parca-Listeleri.xlsx"))
    write_html(all_critical, os.path.join(FOLDER, "kritik-parca-listeleri.html"))

    summary = {g: len(all_critical[g]) for g in all_critical}
    print("OK", summary)


if __name__ == "__main__":
    main()
