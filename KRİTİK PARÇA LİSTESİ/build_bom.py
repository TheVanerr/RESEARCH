# -*- coding: utf-8 -*-
"""Ürün ağaçlarından makine bazlı BOM sayfası (urun-agaci-bom.html) — sınıflandırma yok, yalnızca BOM."""
from __future__ import annotations

import base64
import io
import json
import os
import re

import openpyxl
from PIL import Image

import build_kritik_parca as bkp
import build_kritik_sablon as bks

FOLDER = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(FOLDER, "urun-agaci-bom.html")
GROUP_ORDER = ("LYM", "KBN", "VDL")
# Müşteri BOM'u: yalnızca makinenin fonksiyonunu etkileyen komponentler listelenir.
# Önce HARIC (kesin çıkar), sonra FONKSIYONEL (beyaz liste) uygulanır; ikisine de uymayan
# (sac, gövde, şasi, borulama, bağlantı elemanı, klemens, tesisat parçası vb.) gösterilmez.
HARIC = [
    ("LEO pompa", r"^POMPA LEO(?=\s|$)"),
    ("etiket", r"ETİKET"),
    ("ambalaj", r"^(AMBALAJ|STREÇ FİLM|FOAMBOARD)"),
    ("sac / yapısal", r"SACI?(?=\s|$)|\bSAC\b|DELRİNİ(?!.*TEKER)|TERMOSTAT ADAPTÖRÜ|ELEKTRİK PANO BAĞLANTI|DURDURMA APARATI|SWITCH KOLU"),
]
# Elle çıkarılan stok kodları (tüm makinelerin BOM'larından).
HARIC_KOD = {
    "10 03672 01",  # PYM 1 ELEKTRİK KUTUSU KOMPLESİ (PYM950-1150-1350)
    "07 08286",     # REDÜKTÖR VE MOTOR KOMP SEW LYM 950-1150
    "07 05248",     # LYM EXTRA KAZAN ISITICI KOMPLESİ
    "07 08290",     # REDÜKTÖR VE MOTOR KOMP SEW LYM-KBN 1550
    "10 03672 03",  # PYM 2 ELEKTRİK KUTUSU KOMPLESİ (PYM1400-1550)
    "07 08288",     # REDÜKTÖR VE MOTOR KOMP SEW LYM 1350-1400
    "07 05389",     # MANUEL YIKAMA TABANCASI KOMPLESİ KBN
    "07 08367",     # REDÜKTÖR VE MOTOR KOMP SEW KBN 1650
    "07 06521",     # YAĞ AYIRICI ÜNİTESİ SANDIKLI
    "07 15387",     # KBN 1B 1650 OTOMATİK YÜKLEME KOMPLESİ
    "07 15391",     # KBN 1B 1850 OTOMATİK YÜKLEME KOMPLESİ
    "07 08381",     # REDÜKTÖR VE MOTOR KOMP SEW KBN 1850
    "07 08386",     # REDÜKTÖR VE MOTOR KOMP SEW KBN 2050
    "10 03672 26",  # KBN 2B SIEMENS'Lİ ELEKTRİK KUTUSU KOMPLESİ
    "07 00511",     # KBN DOZAJ ÜNİTESİ 1 BANYO İÇİN
    "07 05388",     # FİLTRE DOLULUK SİSTEM KOMPLESİ KBN
    "02 00905",     # VBR TİP 3
    "02 01801",     # VDL 40-50 ASANSÖRLÜ PARÇA YÜKLEME KONVEYÖRLÜ
    "07 09689",     # VBR TİP 3 SANDIKLI
    "07 00512",     # VDL DOZAJ ÜNİTESİ
    "02 01803",     # VDL 40-50 ASANSÖRLÜ PARÇA YÜKLEME KONVEYÖR SANDIKLI
    "02 02101",     # KASA DEVİRME TİP 1
    "02 02102",     # KASA DEVİRME TİP 2
    "02 00902",     # VİBRASYONLU PARÇA YÜKLEME TİP 1
    "07 09677",     # VİBRASYONLU PARÇA YÜKLEME TİP 1 (SANDIKLI)
    "07 14658",     # SU BASINÇ ÖLÇME KOMPLESİ (alt kırılımlardan da çıkar)
}
# Makine bazlı elle düzeltmeler: makine -> {stok kodu: (ad, adet, birim) | None}. None = kalemi çıkar;
# kayıt varsa adet/ad güncellenir, kalem yoksa eklenir (yerine geçtiği kalemin sırasına).
REZ_8000 = "REZİSTANS KOMPLESİ 8000W 50CM DÜZ DİKİŞSİZ"
DUZELTME = {
    "KBN 2B 1650": {"07 00309": None, "07 15142": (REZ_8000, 4, "adet")},  # 2 tank x 2
    "KBN 2B 1850": {"07 00309": None, "07 15142": (REZ_8000, 4, "adet")},  # 2 tank x 2
    "KBN 2B 2050": {"07 00309": None, "07 15142": (REZ_8000, 6, "adet")},  # 2 tank x 3
}
# Grup bazlı ekler: gruptaki her makineye uygulanır (aynı biçim).
EMNIYET_G9SB = ("EMNİYET RÖLESİ G9SB2002AACDC241 OMRON G9SX", 1, "adet")
GRUP_EK = {g: {"10 08330": EMNIYET_G9SB} for g in ("LYM", "KBN", "VDL")}
FONKSIYONEL = [
    ("pompa", r"^POMPA |DOZAJ POMPASI"),
    ("motor / tahrik", r"^REDÜKTÖR (VE MOTOR|MOTORU|EN:)|ZİNCİR|YATAKLAMA TEKERİ|SEPET TEKER DELRİNİ|AMORTİSÖR"),
    ("ısıtma", r"^REZİSTANS|KAZAN ISITICI"),
    ("sensör / emniyet", r"SENSÖR(?!.*KABLOSU)|REFLEKTÖR|IŞIK BARİYERİ|TERMOKUPU?L|SEVİYE BEKÇİSİ|BASINÇ ŞALTERİ|^SWITCH|INTERLOCK|^EL KORUMA (FOTOSELLİ|LAZER)"),
    ("elektrik", r"TERMİK RÖLESİ|KONTAKTÖR|RÖLE|KONTAK BLOK|^BUTON|LAMBASI|ZAMANLAYICI|ELEKTRİK KUTUSU KOMP|ELEKTRİK KUT\.KOMP|^DÜZ FİŞ|^DUVAR PRİZİ|^MAKİNA PRİZİ"),
    ("valf / nozzle", r"P\.VALF|VANA(?!.*KAPAMA)|^NOZZLE"),
    ("fonksiyon ünitesi", r"FİLTRE|YAĞ SIYIRICI KOMPLESİ|YAĞ AYIRICI|KURUTMA|BUHAR|OTOMATİK (TANK|BOŞALTMA|YÜKLEME)|DOZAJ ÜNİTESİ|UYARI SİSTEMİ|KASKAT|HAVA BASINÇ ÖLÇME|AYDINLATMA|YIKAMA TABANCASI|PARÇA YÜKLEME|^VBR|KASA DEVİRME|SALYANGOZLU FAN"),
]


# "Normal" görünüm: kompleler açılır, yalnızca elektrik bağlantılı ürünler ve pano içi malzemeler kalır.
ELEKTRIKLI = (r"MOTOR|(?<!DİYAFRAMLI )POMPA|\bFAN\b|BLOWER|YAĞ BUHARI TOPLAMA|REZİSTANS|ISITICI|SENSÖR(?!.*KABLOSU)|"
              r"TERMOKUPU?L|TERMOSTAT|SEVİYE BEKÇİSİ|PASLANMAZ ŞAMANDIRA|SELENOİD|ŞALTER|SWITCH|INTERLOCK|IŞIK BARİYERİ|P\.VALF|BOBİN|SİGORTA|R[ÖO]LE|"
              r"KONTAKTÖR|KONTAK|^BUTON|LAMBA|GÜÇ KAYNAĞI|ORDEL|PLC|OPERATÖR PANELİ|NB7W|İNVERTÖR|^TMŞ|NSX|PRİZ|FİŞ|"
              r"ZAMANLAYICI|ELEKTRİK KUT")
# Alt kırılımı girilmemiş ama elektrik içermesi beklenmeyen kompleler (uyarı listesine alınmaz).
MEKANIK_KOMPLE = r"FİLTRE|NOZZLE|HAVA BIÇAĞI|ŞAMANDIRA|YATAKLAMA|KKO|KURUTMA ODASI"


def flatten(items: list[dict], subs: dict[str, list[dict]]) -> tuple[list[dict], list[dict]]:
    """Kalemleri alt kırılımlara kadar açar; her elektrikli malzeme bir kez (kaynak kompleleriyle).
    İkinci liste: alt kırılımı girilmemiş, elektrikli içeriği bilinmeyen kompleler."""
    out: dict[str, dict] = {}
    unknown: dict[str, dict] = {}

    def walk(c: str, n: str, p: str, src: str, seen: frozenset) -> None:
        if subs.get(c) and c not in seen:
            for s in subs[c]:
                walk(s["c"], s["n"], s["p"], src or n, seen | {c})
            return
        name = _norm(n)
        if not re.search(ELEKTRIKLI, name):
            if c[:2] in ("02", "07") and not re.search(MEKANIK_KOMPLE, name):
                unknown.setdefault(c, {"c": c, "n": n, "p": p, "s": [src] if src else []})
            return
        cur = out.setdefault(c, {"c": c, "n": n, "p": p, "s": []})
        if src and src not in cur["s"]:
            cur["s"].append(src)

    for it in items:
        walk(it["c"], it["n"], it["p"], "", frozenset())
    return list(out.values()), list(unknown.values())


# Komple fotoğrafı için ana bileşen önceliği (alt kırılımda ilk eşleşen ve fotoğrafı olan seçilir).
ANA_BILESEN = [
    r"POMPA ",
    r"IŞIK BARİYERİ|^SENSÖR|SEVİYE SENSÖRÜ|AKTÜATÖR ",
    r"FAN|BLOWER|YAĞ BUHARI",
    r"REDÜKTÖR",
    r"^REZİSTANS",
    r"FİLTRE",
    r"OPERATÖR PANELİ|NB7W|DİJ\.PAN|PLC",  # elektrik kutuları -> panel / PLC
    r"INTERLOCK|^SWITCH",
    r"TERMOSTAT",
    r"ZAMAN",
    r"RÖLE",
    r"KONTAKTÖR",
]


def _norm(name: str) -> str:
    return str(name).strip().replace("i", "İ").upper()


def haric_mi(name: str) -> str | None:
    """Gösterilmeyecekse sebebini döndürür; fonksiyonel komponentse None."""
    n = _norm(name)
    hit = next((lbl for lbl, pat in HARIC if re.search(pat, n)), None)
    if hit:
        return hit
    return None if any(re.search(pat, n) for _, pat in FONKSIYONEL) else "fonksiyonel değil"


def dedupe(items: list[dict]) -> list[dict]:
    """Aynı kod bir makinede birden çok satırdaysa tek satır (ilk sıra, en büyük miktar)."""
    out: dict[str, dict] = {}
    for it in items:
        cur = out.get(it["code"])
        if cur is None:
            out[it["code"]] = dict(it)
            continue
        try:
            if float(it["qty"] or 0) > float(cur["qty"] or 0):
                cur["qty"] = it["qty"]
        except (TypeError, ValueError):
            pass
    return list(out.values())


UNITS = {"AD": "adet", "ADET": "adet", "HY": "adet",  # HY: ERP birim kodu (komple), adet olarak gösterilir
          "MT": "m", "M": "m", "KG": "kg", "LT": "lt", "M2": "m²", "TK": "takım"}


def parse_kbn(path: str) -> list[tuple[str, list[dict]]]:
    """KBN geniş tablo: her makine bloğu 'Kaynak Kodu, Kaynak Adı, Miktar, Birim'."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    blocks = [(i, h.strip()) for i, h in enumerate(rows[0]) if isinstance(h, str) and "KBN" in h.upper()]
    out = []
    for col, vname in blocks:
        items = []
        for row in rows[3:]:
            if col + 5 >= len(row):
                continue
            code, name, qty, unit = bkp.norm_code(row[col + 2]), row[col + 3], row[col + 4], row[col + 5]
            if code and name:
                items.append({"code": code, "name": str(name).strip(), "qty": qty,
                              "unit": UNITS.get(str(unit or "").strip().upper(), str(unit or "").strip().lower())})
        out.append((re.sub(r"\s*PARÇA YIKAMA MAKİNESİ\s*$", "", vname), items))
    return out


def load_machines() -> dict[str, list[tuple[str, list[dict]]]]:
    files = os.listdir(FOLDER)
    machine = next(f for f in files if f.endswith(".xlsx") and "MAK" in f.upper())
    wb = openpyxl.load_workbook(os.path.join(FOLDER, machine), read_only=True, data_only=True)
    groups: dict[str, list[tuple[str, list[dict]]]] = {}
    for g in ("LYM", "VDL"):
        groups[g] = [(v, [{**it, "unit": ""} for it in items])
                     for v, items in bkp.parse_variant_columns(wb[g]).items()]
    wb.close()
    kbn = []
    for tag in ("1B", "2B"):
        f = next(f for f in files if tag in f and "KBN" in f.upper() and f.endswith(".xlsx"))
        kbn += parse_kbn(os.path.join(FOLDER, f))
    groups["KBN"] = kbn
    for g, L in groups.items():
        for v, items in L:
            apply_fix(items, {**GRUP_EK.get(g, {}), **DUZELTME.get(v, {})})
    return groups


def apply_fix(items: list[dict], fix: dict) -> None:
    pos = next((i for i, it in enumerate(items) if it["code"] in fix), len(items))
    items[:] = [it for it in items if fix.get(it["code"], ...) is not None]
    for code, val in fix.items():
        if val is None:
            continue
        name, qty, unit = val
        same = [it for it in items if it["code"] == code]
        for it in same:
            it.update(name=name, qty=qty, unit=unit)
        if not same:
            items.insert(min(pos, len(items)), {"code": code, "name": name, "qty": qty, "unit": unit})


def load_subparts() -> dict[str, list[dict]]:
    path = os.path.join(FOLDER, "bom-alt-kirilimlar.json")
    if not os.path.isfile(path):
        return {}
    with open(path, encoding="utf-8") as f:
        raw = json.load(f).get("kompleler", {})
    return {k: [{"c": a["kod"], "n": a["ad"]} for a in v.get("alt_malzemeler", []) if a["kod"] not in HARIC_KOD]
            for k, v in raw.items() if k not in HARIC_KOD}


def load_full_names() -> dict[str, dict[str, str]]:
    """'Ürün tam adı' sütunu: stok kodu -> {tr, en, de} (bom-urun-adlari.json)."""
    path = os.path.join(FOLDER, "bom-urun-adlari.json")
    if not os.path.isfile(path):
        return {}
    with open(path, encoding="utf-8") as f:
        return json.load(f).get("adlar", {})


# Yalnızca BOM sayfası için ek fotoğraf aileleri: (ad deseni, referans stok kodu)
BOM_FOTO_AILELERI = [
    (r"BLOWER KURUTMA", "10 01524"),  # blower ile kurutma kompleleri -> tek kademeli blower
    (r"KÖRÜKLÜ FAN ERF", "07 17295"),  # ERF fan ailesi (aynı görünüm)
    (r"POMPA LEO AMS", "10 00695"),     # LEO AMS serisi paslanmaz santrifüj pompalar (aynı gövde)
    (r"PASLANMAZ ŞAMANDIRA", "10 00296"),  # paslanmaz şamandıralı seviye şalteri
    (r"MOTORLU .*MR\d{3}", "10 01718"),  # Yılmaz MR serisi ayaklı helisel redüktörlü motorlar
    # Aynı seri / aynı gövde -> klasördeki kardeş ürünün fotoğrafı
    (r"^SİGORTA A9F741", "10 01421"),        # Acti9 iC60N 1 kutuplu
    (r"KAÇAK AKIMLI A9R", "10 00303"),       # Acti9 iID 4 kutuplu kaçak akım
    (r"^SİNYAL LAMBASI S222|^SİNYAL LAMBA SARI 22MM", "10 02780"),  # Emas S2 22 mm sinyal lambası
    (r"B7 B MAVİ LED", "10 00275"),          # Emas B serisi mavi LED lamba
    (r"^İNVERTÖR VFD0\d\dEL", "10 16319"),   # Delta VFD-EL gövdesi
    (r"CT-2467", "10 03253"),                # pano LED bant armatür
    (r"LRS 350/24", "10 04903"),             # Mean Well LRS-350-24
    (r"^SWITCH BS10", "10 16235"),           # BS10xx mini sınır şalteri
    (r"^REDÜKTÖR EN:", "10 01002"),          # Yılmaz EN sonsuz vidalı redüktör
    (r"^KÜRESEL VANA .*GALVANİZLİ", "10 00543"),
    (r"^KÜRESEL VANA", "10 00544"),
    (r"^TORBA FİLTRE", "10 05378"),
    (r"^TAŞ FİLTRE", "10 00470"),
    (r"HASSAS FİLTRE KOMPLESİ|HASSAS FİLTRASYON .*YUVASI", "10 05378"),  # torba filtreli hassas filtre
    (r"OTOMATİK TANK DOLUM", "07 16791"),    # dolum ünitesi -> seviye sensörü
    (r"KURUTMA KOMPLESİ", "07 17295"),       # kurutma -> fan
]


def photo_for(code: str, name: str, photos: dict[str, str]) -> str:
    """Önce ürünün kendi stok kodlu fotoğrafı, yoksa aile fotoğrafı."""
    own = photos.get(bks.norm_siparis(code) or "", "")
    if own:
        return own
    return next((photos.get(ref) for pat, ref in bks.FOTO_AILELERI + BOM_FOTO_AILELERI if re.search(pat, name)), None) or ""


def embed(rel: str, max_side: int = 640) -> str:
    with Image.open(os.path.join(FOLDER, rel)) as im:
        im = im.convert("RGBA")
        bg = Image.new("RGB", im.size, "white")
        bg.paste(im, mask=im.split()[3])
        bg.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        bg.save(buf, "JPEG", quality=80, optimize=True)
    return "data:image/jpeg;base64," + base64.b64encode(buf.getvalue()).decode("ascii")


def fmt_qty(q) -> str:
    if q is None or q == "":
        return ""
    try:
        f = float(q)
    except (TypeError, ValueError):
        return str(q)
    return str(int(f)) if f.is_integer() else f"{f:.3f}".rstrip("0").rstrip(".").replace(".", ",")


def build_data() -> tuple[dict, dict]:
    machines, subs, photos = load_machines(), load_subparts(), bks.photo_map()
    imgs: dict[str, str] = {}

    def pic(code: str, name: str) -> str:
        rel = photo_for(code, name, photos)
        if rel and rel not in imgs and os.path.isfile(os.path.join(FOLDER, rel)):
            imgs[rel] = embed(rel)
        return rel if rel in imgs else ""

    data = {}
    for g in GROUP_ORDER:
        data[g] = []
        for v, items in machines[g]:
            kept = [it for it in dedupe(items) if it["code"] not in HARIC_KOD and not haric_mi(it["name"])]
            data[g].append({"m": v, "x": len(items) - len(kept),
                            "items": [{"c": it["code"], "n": it["name"], "q": fmt_qty(it["qty"]), "u": it["unit"],
                                       "p": pic(it["code"], it["name"])} for it in kept]})
    # Alt malzemeler: kod -> [{c, n, p}] (iç içe kompleler sayfada açılır)
    sub_data = {k: [{**s, "p": pic(s["c"], s["n"])} for s in v] for k, v in subs.items()}  # elle seçilmiş, filtre yok
    # Fotoğrafı olmayan komple, alt kırılımındaki ana bileşenin fotoğrafını kullanır (ör. boşaltma ünitesi -> pompa).
    komple_foto: dict[str, str] = {}

    def main_photo(code: str, seen: frozenset = frozenset()) -> str:
        if code in komple_foto or code in seen:
            return komple_foto.get(code, "")
        cands = [(s, s["p"] or (main_photo(s["c"], seen | {code}) if s["c"] in sub_data else "")) for s in sub_data.get(code, [])]
        cands = [(s, p) for s, p in cands if p]
        rank = lambda s: next((i for i, pat in enumerate(ANA_BILESEN) if re.search(pat, _norm(s["n"]))), len(ANA_BILESEN))
        komple_foto[code] = min(cands, key=lambda sp: rank(sp[0]))[1] if cands else ""
        return komple_foto[code]

    for code in sub_data:
        main_photo(code)
    for g in data.values():
        for m in g:
            for it in m["items"]:
                it["p"] = it["p"] or komple_foto.get(it["c"], "")
    for v in sub_data.values():
        for s in v:
            s["p"] = s["p"] or komple_foto.get(s["c"], "")
    # Normal görünüm: gruptaki tüm makinelerin elektrikli malzemeleri tek BOM, her malzeme bir kez.
    flat = {}
    for g, ms in data.items():
        items, unk = flatten([it for m in ms for it in m["items"]], sub_data)
        flat[g] = {"items": items, "unk": unk}
    names = load_full_names()
    shown = {it["c"] for g in data.values() for m in g for it in m["items"]} | {x["c"] for v in sub_data.values() for x in v}
    full = {c: names[c] for c in sorted(shown) if c in names}
    missing = sorted(shown - set(full))
    return {"groups": data, "flat": flat, "subs": sub_data, "full": full, "missing": missing}, imgs


HTML_TEMPLATE = r'''<!DOCTYPE html>
<html lang="tr">
<head>
  <meta charset="utf-8"/>
  <meta name="viewport" content="width=device-width, initial-scale=1"/>
  <title>Ürün Ağaçları</title>
  <style>
    :root {
      --bg:#f4f5f7; --surface:#ffffff; --surface-2:#f0f2f5; --line:#e4e7ec;
      --text:#101828; --muted:#667085; --faint:#98a2b3;
      --accent:#2f5bea; --accent-soft:#eaf0ff; --warn:#b54708; --warn-soft:#fef0c7;
      --shadow:0 1px 2px rgba(16,24,40,.05), 0 1px 3px rgba(16,24,40,.06);
      --shadow-lg:0 12px 32px rgba(16,24,40,.14);
      color-scheme:light;
    }
    @media (prefers-color-scheme: dark) {
      :root:not([data-theme="light"]) {
        --bg:#0c111d; --surface:#161b26; --surface-2:#1f242f; --line:#2a313e;
        --text:#f5f5f6; --muted:#94969c; --faint:#61646c; --accent:#7b9bff; --accent-soft:#1c2a52; --warn:#fdb022; --warn-soft:#4e2a0a;
        --shadow:0 1px 2px rgba(0,0,0,.4); --shadow-lg:0 12px 32px rgba(0,0,0,.5); color-scheme:dark;
      }
    }
    :root[data-theme="dark"] {
      --bg:#0c111d; --surface:#161b26; --surface-2:#1f242f; --line:#2a313e;
      --text:#f5f5f6; --muted:#94969c; --faint:#61646c; --accent:#7b9bff; --accent-soft:#1c2a52; --warn:#fdb022; --warn-soft:#4e2a0a;
      --shadow:0 1px 2px rgba(0,0,0,.4); --shadow-lg:0 12px 32px rgba(0,0,0,.5); color-scheme:dark;
    }
    * { box-sizing:border-box; }
    html, body { margin:0; }
    body { background:var(--bg); color:var(--text); font:14px/1.45 "Segoe UI Variable Text","Segoe UI",system-ui,-apple-system,sans-serif; -webkit-font-smoothing:antialiased; }
    button { font:inherit; color:inherit; cursor:pointer; }
    .mono { font-family:"Cascadia Mono",Consolas,ui-monospace,monospace; }

    .top { position:sticky; top:0; z-index:20; display:flex; align-items:center; gap:12px 16px; flex-wrap:wrap;
      padding:12px 24px; background:color-mix(in srgb, var(--surface) 88%, transparent); backdrop-filter:blur(10px); border-bottom:1px solid var(--line); }
    .brand { display:flex; align-items:center; gap:10px; margin-right:8px; }
    .logo { width:34px; height:34px; border-radius:10px; display:grid; place-items:center; background:var(--accent); color:#fff; font-weight:800; font-size:12px; }
    .brand h1 { margin:0; font-size:15px; font-weight:700; letter-spacing:-.01em; }
    .brand small { display:block; color:var(--muted); font-size:12px; }
    .seg { display:inline-flex; padding:3px; background:var(--surface-2); border-radius:10px; gap:2px; flex-wrap:wrap; }
    .seg button { border:0; background:transparent; padding:6px 14px; border-radius:8px; font-weight:600; color:var(--muted); display:inline-flex; align-items:center; gap:6px; white-space:nowrap; }
    .seg button .n { font-size:11px; font-weight:700; color:var(--faint); }
    .seg button.on { background:var(--surface); color:var(--text); box-shadow:var(--shadow); }
    .seg button.on .n { color:var(--accent); }
    .search { flex:1; min-width:180px; max-width:360px; margin-left:auto; position:relative; }
    .search svg { position:absolute; left:11px; top:50%; transform:translateY(-50%); color:var(--faint); }
    .search input { width:100%; height:36px; padding:0 12px 0 34px; border-radius:10px; border:1px solid var(--line); background:var(--surface); color:var(--text); font:inherit; }
    .search input:focus { outline:none; border-color:var(--accent); box-shadow:0 0 0 3px var(--accent-soft); }
    .tools { display:flex; gap:8px; }
    .icon-btn { height:36px; min-width:36px; padding:0 10px; border-radius:10px; border:1px solid var(--line); background:var(--surface); display:inline-flex; align-items:center; justify-content:center; gap:6px; color:var(--muted); font-size:13px; font-weight:600; }
    .icon-btn:hover { color:var(--text); border-color:var(--faint); }

    main { padding:20px 24px 40px; max-width:1400px; margin:0 auto; }
    .machines { margin-bottom:16px; }
    .machines .seg { background:transparent; padding:0; gap:8px; }
    .machines .seg button { border:1px solid var(--line); background:var(--surface); padding:8px 14px; border-radius:10px; }
    .machines .seg button.on { border-color:var(--accent); color:var(--accent); box-shadow:0 0 0 3px var(--accent-soft); }

    .head { display:flex; align-items:flex-end; justify-content:space-between; gap:12px; flex-wrap:wrap; margin-bottom:12px; }
    .head h2 { margin:0; font-size:20px; letter-spacing:-.01em; }
    .stats { display:flex; gap:18px; color:var(--muted); font-size:13px; }
    .stats b { color:var(--text); font-size:15px; font-variant-numeric:tabular-nums; margin-right:4px; }

    .bom { background:var(--surface); border:1px solid var(--line); border-radius:14px; overflow:auto; box-shadow:var(--shadow); }
    table { width:100%; border-collapse:collapse; }
    th { position:sticky; top:0; background:var(--surface-2); color:var(--muted); font-size:11px; font-weight:700; text-transform:uppercase; letter-spacing:.05em; text-align:left; padding:10px 12px; white-space:nowrap; z-index:1; }
    td { padding:7px 12px; border-top:1px solid var(--line); vertical-align:middle; }
    td.no { color:var(--faint); font-size:12px; width:44px; text-align:right; font-variant-numeric:tabular-nums; }
    td.ph { width:96px; padding:8px 10px; }
    .thumb { width:76px; height:76px; border-radius:10px; object-fit:contain; background:#fff; display:block; cursor:zoom-in; border:1px solid var(--line); }
    .thumb.none { display:grid; place-items:center; color:var(--faint); cursor:default; border-style:dashed; background:var(--surface-2); }
    td.code { white-space:nowrap; color:var(--accent); font-size:12.5px; font-weight:600; }
    td.qty { text-align:right; font-weight:700; font-variant-numeric:tabular-nums; white-space:nowrap; width:80px; }
    td.unit { color:var(--muted); font-size:12px; width:70px; }
    tr:hover td { background:var(--surface-2); }
    .tg { border:0; background:var(--accent-soft); color:var(--accent); width:22px; height:22px; border-radius:6px; display:inline-grid; place-items:center; margin-right:8px; vertical-align:middle; transition:transform .15s; }
    .tg.open { transform:rotate(90deg); }
    .tag { font-size:11px; font-weight:600; color:var(--accent); background:var(--accent-soft); border-radius:6px; padding:1px 7px; margin-left:8px; white-space:nowrap; }
    tr.sub td { background:color-mix(in srgb, var(--surface-2) 55%, var(--surface)); font-size:13px; }
    tr.sub td.name { padding-left:calc(12px + var(--lvl, 1) * 30px); }
    tr.sub td.name::before { content:"└"; color:var(--faint); margin-right:8px; }
    tr.sub .thumb { width:56px; height:56px; }
    .thumb:not(.none):hover { transform:scale(1.06); box-shadow:var(--shadow-lg); }
    .thumb { transition:transform .12s ease, box-shadow .12s ease; }
    mark { background:#fde68a; color:#1f2937; border-radius:3px; padding:0 1px; }
    .empty { text-align:center; padding:48px 20px; color:var(--muted); }

    dialog { border:0; padding:0; background:transparent; max-width:min(560px, calc(100% - 32px)); }
    dialog::backdrop { background:rgba(12,17,29,.6); }
    .lb { background:#fff; border-radius:16px; overflow:hidden; box-shadow:var(--shadow-lg); }
    .lb img { display:block; width:100%; height:auto; max-height:70vh; object-fit:contain; background:#fff; }
    .lb .cap { padding:12px 16px; background:var(--surface); color:var(--text); font-size:13px; }
    .lb .cap b { display:block; font-size:14px; }
    .lb .capfull { color:var(--muted); margin-top:4px; }
    td.full { color:var(--muted); font-size:12.5px; min-width:220px; }
    td.full.none { color:var(--faint); }
    th.full { min-width:220px; }
    .var { font-size:11px; font-weight:600; color:var(--warn); background:var(--warn-soft); border-radius:6px; padding:1px 7px; margin-left:8px; white-space:nowrap; cursor:help; }
    td.src { color:var(--muted); font-size:12px; min-width:160px; max-width:280px; }
    td.where { color:var(--muted); font-size:12.5px; }
    .note { color:var(--muted); font-size:12.5px; margin:-4px 0 12px; }
    .icon-btn.on { color:var(--accent); border-color:var(--accent); background:var(--accent-soft); }
    .stats .icon-btn { height:28px; font-size:12px; }
    .stats { align-items:center; flex-wrap:wrap; }
    .sec { margin-top:28px; }
    .sec h3 { margin:0 0 4px; font-size:15px; display:flex; align-items:center; gap:8px; }
    .sec h3 .cnt { font-size:12px; font-weight:700; color:var(--warn); background:var(--warn-soft); border-radius:999px; padding:1px 8px; }
    .sec p { margin:0 0 10px; color:var(--muted); font-size:13px; }
__LANGCSS__

    @media (max-width:640px) {
      .top { padding:10px 16px; }
      .search { order:5; max-width:none; flex-basis:100%; }
      .tools { margin-left:auto; }
      main { padding:16px 16px 32px; }
      td.ph { width:72px; } .thumb { width:56px; height:56px; } tr.sub .thumb { width:44px; height:44px; }
      td.no, th.no, td.unit, th.unit { display:none; }
      .tools .lbl { display:none; }
    }
    @media print {
      .top .search, .tools, .machines { display:none !important; }
      .top { position:static; backdrop-filter:none; }
      body { background:#fff; }
      .bom { box-shadow:none; border:0; }
      th { position:static; }
    }
  </style>
</head>
<body>
  <header class="top">
    <div class="brand">
      <div class="logo">BOM</div>
      <div><h1 data-i="title"></h1><small data-i="subtitle"></small></div>
    </div>
    <nav class="seg" id="view"></nav>
    <nav class="seg" id="groups"></nav>
    <label class="search">
      <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round"><circle cx="11" cy="11" r="7"/><path d="m20 20-3.5-3.5"/></svg>
      <input type="search" id="q" autocomplete="off"/>
    </label>
    <div class="tools">
      <div class="lang" id="lang">
        <button class="lang-btn" id="langBtn" aria-haspopup="listbox" aria-expanded="false"></button>
        <ul class="lang-menu" id="langMenu" role="listbox" tabindex="-1" hidden></ul>
      </div>
      <button class="icon-btn" id="csvBtn">
        <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round"><path d="M12 3v12m0 0-4-4m4 4 4-4M5 21h14"/></svg><span class="lbl">CSV</span>
      </button>
      <button class="icon-btn" id="printBtn">
        <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linejoin="round"><path d="M7 8V3h10v5M7 17H4v-7h16v7h-3M7 14h10v7H7z"/></svg>
      </button>
      <button class="icon-btn" id="themeBtn">
        <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2"><circle cx="12" cy="12" r="9"/><path d="M12 3a9 9 0 0 1 0 18z" fill="currentColor"/></svg>
      </button>
    </div>
  </header>

  <main>
    <div class="machines"><div class="seg" id="machines"></div></div>
    <div class="head">
      <h2 id="title"></h2>
      <div class="stats" id="stats"></div>
    </div>
    <div class="note" id="note"></div>
    <div class="bom" id="bom"></div>
    <div id="extra"></div>
  </main>

  <dialog id="lb"><div class="lb"><img id="lbImg" alt=""/><div class="cap" id="lbCap"></div></div></dialog>

  <script>
    const DATA = __DATA__;
    const IMGS = __IMGS__;
    __FLAGS__
    const LANGS = [
      { id:'tr', code:'TR', name:'Türkçe', locale:'tr' },
      { id:'en', code:'EN', name:'English', locale:'en' },
      { id:'de', code:'DE', name:'Deutsch', locale:'de' }
    ];
    const UI = {
      tr: { title:'Ürün Ağaçları', subtitle:'Makine bazlı malzeme listesi', search:'Stok kodu veya ürün adı ara',
        csv:'Seçili makinenin BOM\'unu CSV olarak indir', print:'Yazdır', theme:'Temayı değiştir', lang:'Dil',
        groups:'Makine grubu', machines:'Makine', items:'kalem', komple:'alt malzemeli komple', photos:'fotoğraflı',
        hidden:'gizli kalem', matches:'eşleşme', empty:'Eşleşen malzeme yok.', hCode:'Stok kodu', hName:'Stok adı',
        hFull:'Ürün tam adı', hQty:'Adet', hUnit:'Birim', sub:'alt', subAria:'Alt malzemeler', unknown:'Tam ad henüz girilmedi',
        csvNo:'Sıra', csvParent:'Üst komple',
        view:'Görünüm', vFlat:'Normal', vTree:'Alt kırılımlı', flatItems:'elektrikli malzeme', hSrc:'Kaynak komple', direct:'Makinede doğrudan',
        onlyHere:'yalnız bu makinede', inSeries:'makinede', varied:'seride farklı', diffOnly:'Yalnız farklılar', series:'Seri', machinesN:'makine',
        missTitle:'Serideki diğer makinelerde olup bu makinede olmayanlar', missNote:'Opsiyon farkı mı, yoksa eksik ürün ağacı verisi mi kontrol edin.',
        missing:'seride var, burada yok', hWhere:'Bulunduğu makineler', unkTitle:'Alt kırılımı girilmemiş kompleler', unk:'alt kırılımı eksik komple',
        unkNote:'Bu komplelerin içindeki elektrikli malzemeler listeye alınamadı. ERP ürün ağacı eklenirse otomatik listelenir.',
        flatNote:'Gruptaki tüm makinelerin ürün ağaçları birleştirildi, kompleler açıldı: yalnızca motor, pompa, fan, rezistans, sensör, valf bobini ve pano içi elektrik malzemeleri. Her malzeme bir kez yazılır.', flatScope:'Kapsanan makineler',
        excluded:'Gösterilmeyenler: tekrar eden satırlar, LEO pompalar, etiketler, sac / gövde / şasi / borulama, bağlantı elemanları, kablo, rakor, klemens ve diğer fonksiyonel olmayan kalemler',
        units:{} },
      en: { title:'Bills of Materials', subtitle:'Parts list per machine', search:'Search stock code or name',
        csv:'Download the selected machine\'s BOM as CSV', print:'Print', theme:'Toggle theme', lang:'Language',
        groups:'Machine group', machines:'Machine', items:'items', komple:'assemblies with sub-parts', photos:'with photo',
        hidden:'hidden items', matches:'matches', empty:'No matching parts.', hCode:'Stock code', hName:'Stock name',
        hFull:'Full product name', hQty:'Qty', hUnit:'Unit', sub:'sub', subAria:'Sub-parts', unknown:'Full name not entered yet',
        csvNo:'No.', csvParent:'Parent assembly',
        view:'View', vFlat:'Normal', vTree:'With sub-parts', flatItems:'electrical parts', hSrc:'Source assembly', direct:'Directly on machine',
        onlyHere:'only on this machine', inSeries:'machines', varied:'differ in series', diffOnly:'Differences only', series:'Series', machinesN:'machines',
        missTitle:'On other machines of the series but not on this one', missNote:'Check whether this is an option difference or missing BOM data.',
        missing:'in series, not here', hWhere:'Found on', unkTitle:'Assemblies without sub-parts entered', unk:'assemblies without sub-parts',
        unkNote:'Electrical parts inside these assemblies could not be listed. They appear automatically once the ERP BOM is added.',
        flatNote:'BOMs of all machines in the group merged, assemblies expanded: only motors, pumps, fans, heaters, sensors, valve coils and control-cabinet electrical parts. Each part listed once.', flatScope:'Machines covered',
        excluded:'Not shown: duplicate rows, LEO pumps, labels, sheet metal / body / frame / piping, fasteners, cables, cable glands, terminals and other non-functional items',
        units:{ adet:'pcs', 'takım':'set' } },
      de: { title:'Stücklisten', subtitle:'Materialliste je Maschine', search:'Lagercode oder Bezeichnung suchen',
        csv:'Stückliste der gewählten Maschine als CSV herunterladen', print:'Drucken', theme:'Design wechseln', lang:'Sprache',
        groups:'Maschinengruppe', machines:'Maschine', items:'Positionen', komple:'Baugruppen mit Unterteilen', photos:'mit Foto',
        hidden:'ausgeblendete Positionen', matches:'Treffer', empty:'Keine passenden Teile.', hCode:'Lagercode', hName:'Lagerbezeichnung',
        hFull:'Vollständige Produktbezeichnung', hQty:'Menge', hUnit:'Einheit', sub:'Unter', subAria:'Unterteile', unknown:'Vollständige Bezeichnung noch nicht erfasst',
        csvNo:'Nr.', csvParent:'Übergeordnete Baugruppe',
        view:'Ansicht', vFlat:'Normal', vTree:'Mit Unterteilen', flatItems:'elektrische Teile', hSrc:'Herkunftsbaugruppe', direct:'Direkt an der Maschine',
        onlyHere:'nur an dieser Maschine', inSeries:'Maschinen', varied:'in der Serie abweichend', diffOnly:'Nur Abweichungen', series:'Serie', machinesN:'Maschinen',
        missTitle:'An anderen Maschinen der Serie, aber nicht an dieser', missNote:'Prüfen Sie, ob es sich um eine Option oder um fehlende Stücklistendaten handelt.',
        missing:'in der Serie, hier nicht', hWhere:'Vorhanden an', unkTitle:'Baugruppen ohne erfasste Unterteile', unk:'Baugruppen ohne Unterteile',
        unkNote:'Elektrische Teile in diesen Baugruppen konnten nicht aufgelistet werden. Sie erscheinen automatisch, sobald die ERP-Stückliste ergänzt wird.',
        flatNote:'Stücklisten aller Maschinen der Gruppe zusammengeführt, Baugruppen aufgelöst: nur Motoren, Pumpen, Lüfter, Heizungen, Sensoren, Ventilspulen und elektrische Schaltschrankteile. Jedes Teil einmal.', flatScope:'Erfasste Maschinen',
        excluded:'Nicht angezeigt: doppelte Zeilen, LEO-Pumpen, Etiketten, Blech / Gehäuse / Rahmen / Verrohrung, Verbindungselemente, Kabel, Kabelverschraubungen, Klemmen und andere nicht funktionale Positionen',
        units:{ adet:'Stk.', 'takım':'Satz' } }
    };
    const G = Object.keys(DATA.groups);
    const NOIMG = '<svg width="18" height="18" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="1.6" stroke-linejoin="round"><path d="M21 8 12 3 3 8v8l9 5 9-5z"/><path d="m3 8 9 5 9-5M12 13v8"/></svg>';
    const CHEV = '<svg width="12" height="12" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="3" stroke-linecap="round"><path d="m9 6 6 6-6 6"/></svg>';
    const store = {
      get(k, d) { try { return localStorage.getItem('bom.' + k) ?? d; } catch (e) { return d; } },
      set(k, v) { try { localStorage.setItem('bom.' + k, v); } catch (e) {} }
    };
    const st = { g: G.includes(store.get('g')) ? store.get('g') : G[0], m: {}, q: '', open: new Set(),
      lang: UI[store.get('lang')] ? store.get('lang') : 'tr', view: store.get('view') === 'tree' ? 'tree' : 'flat' };
    try { Object.assign(st.m, JSON.parse(store.get('m', '{}'))); } catch (e) {}
    const $ = id => document.getElementById(id);
    const esc = s => String(s ?? '').replace(/&/g,'&amp;').replace(/</g,'&lt;').replace(/>/g,'&gt;').replace(/"/g,'&quot;');
    const low = s => String(s ?? '').toLocaleLowerCase('tr');
    const t = k => UI[st.lang][k] ?? UI.tr[k] ?? k;
    const full = c => ((DATA.full[c] || {})[st.lang]) || '';
    const unit = u => (UI[st.lang].units[u]) || u;
    const machine = () => { const list = DATA.groups[st.g]; return list[Math.min(st.m[st.g] || 0, list.length - 1)]; };
    const hasSub = c => !!(DATA.subs[c] && DATA.subs[c].length);
    const flat = () => st.view === 'flat';
    const srcCell = (it, q) => '<td class="src">' + (it.s.length ? it.s.map(x => hl(x, q)).join('<br>') : esc(t('direct'))) + '</td>';

    function renderFlat(q) {
      const F = DATA.flat[st.g];
      const rows = F.items.map((it, i) => ({ it, i })).filter(({ it }) => matches(it, q) || it.s.some(x => low(x).includes(q)));
      $('title').textContent = st.g + ' · ' + t('vFlat');
      $('stats').innerHTML = '<span><b>' + F.items.length + '</b>' + esc(t('flatItems')) + '</span>' +
        '<span><b>' + F.items.filter(x => x.p).length + '</b>' + esc(t('photos')) + '</span>' +
        '<span><b>' + DATA.groups[st.g].length + '</b>' + esc(t('machinesN')) + '</span>' +
        (F.unk.length ? '<span><b>' + F.unk.length + '</b>' + esc(t('unk')) + '</span>' : '') +
        (q ? '<span><b>' + rows.length + '</b>' + esc(t('matches')) + '</span>' : '');
      $('note').textContent = t('flatNote') + ' ' + t('flatScope') + ': ' + DATA.groups[st.g].map(m => m.m).join(', ') + '.';
      $('bom').innerHTML = !rows.length ? '<div class="empty">' + esc(t('empty')) + '</div>' :
        '<table><thead><tr><th class="no">#</th><th></th><th>' + esc(t('hCode')) + '</th><th>' + esc(t('hName')) +
        '</th><th class="full">' + esc(t('hFull')) + '</th><th>' + esc(t('hSrc')) + '</th><th style="text-align:right">' + esc(t('hQty')) + '</th></tr></thead><tbody>' +
        rows.map(({ it, i }) => '<tr><td class="no">' + (i + 1) + '</td><td class="ph">' + thumb(it.p, it.c, it.n) + '</td>' +
          '<td class="code mono">' + hl(it.c, q) + '</td><td class="name">' + hl(it.n, q) + '</td>' +
          fullCell(it.c, q) + srcCell(it, q) + '<td class="qty">1</td></tr>').join('') + '</tbody></table>';
      $('extra').innerHTML = !F.unk.length ? '' : '<section class="sec"><h3>' + esc(t('unkTitle')) + '<span class="cnt">' + F.unk.length + '</span></h3><p>' +
        esc(t('unkNote')) + '</p><div class="bom"><table><thead><tr><th></th><th>' + esc(t('hCode')) + '</th><th>' + esc(t('hName')) + '</th><th>' +
        esc(t('hSrc')) + '</th></tr></thead><tbody>' + F.unk.map(u => '<tr><td class="ph">' + thumb(u.p, u.c, u.n) + '</td><td class="code mono">' + esc(u.c) +
        '</td><td class="name">' + esc(u.n) + '</td><td class="where">' + esc(u.s.join(', ') || t('direct')) + '</td></tr>').join('') + '</tbody></table></div></section>';
    }

    function hl(text, q) {
      const s = String(text ?? ''); if (!q) return esc(s);
      const i = low(s).indexOf(q); if (i < 0) return esc(s);
      return esc(s.slice(0, i)) + '<mark>' + esc(s.slice(i, i + q.length)) + '</mark>' + esc(s.slice(i + q.length));
    }
    function thumb(p, c, n) {
      return p && IMGS[p] ? '<img class="thumb" src="' + IMGS[p] + '" alt="" loading="lazy" data-c="' + esc(c) + '" data-n="' + esc(n) + '" data-p="' + esc(p) + '"/>'
                          : '<span class="thumb none">' + NOIMG + '</span>';
    }
    function fullCell(c, q) {
      const f = full(c);
      return f ? '<td class="full">' + hl(f, q) + '</td>' : '<td class="full none" title="' + esc(t('unknown')) + '">—</td>';
    }
    const matches = (it, q) => !q || low(it.c).includes(q) || low(it.c).replace(/\s+/g, '').includes(q.replace(/\s+/g, '')) ||
      low(it.n).includes(q) || low(full(it.c)).includes(q);
    function subMatches(c, q, seen = new Set()) {
      if (!q || !hasSub(c) || seen.has(c)) return false; seen.add(c);
      return DATA.subs[c].some(s => matches(s, q) || subMatches(s.c, q, seen));
    }
    const toggle = (key, open, c) => '<button class="tg' + (open ? ' open' : '') + '" data-k="' + esc(key) + '" aria-label="' + esc(t('subAria')) + '">' + CHEV + '</button>';
    const subTag = c => '<span class="tag">' + DATA.subs[c].length + ' ' + esc(t('sub')) + '</span>';

    function subRows(c, q, lvl, path) {
      return DATA.subs[c].map(s => {
        const key = path + '>' + s.c, nested = hasSub(s.c) && !path.split('>').includes(s.c);
        const open = nested && (st.open.has(key) || (q && subMatches(s.c, q)));
        return '<tr class="sub" style="--lvl:' + lvl + '"><td class="no"></td><td class="ph">' + thumb(s.p, s.c, s.n) + '</td>' +
          '<td class="code mono">' + hl(s.c, q) + '</td>' +
          '<td class="name">' + (nested ? toggle(key, open) : '') + hl(s.n, q) + (nested ? subTag(s.c) : '') + '</td>' +
          fullCell(s.c, q) + '<td class="qty"></td><td class="unit"></td></tr>' + (open ? subRows(s.c, q, lvl + 1, key) : '');
      }).join('');
    }

    function renderStatic() {
      document.documentElement.lang = st.lang;
      document.title = t('title');
      document.querySelectorAll('[data-i]').forEach(el => { el.textContent = t(el.dataset.i); });
      $('q').placeholder = t('search');
      $('groups').setAttribute('aria-label', t('groups'));
      $('machines').setAttribute('aria-label', t('machines'));
      const lbl = (el, k) => { el.title = t(k); el.setAttribute('aria-label', t(k)); };
      lbl($('csvBtn'), 'csv'); lbl($('printBtn'), 'print'); lbl($('themeBtn'), 'theme'); lbl($('langBtn'), 'lang');
      const cur = LANGS.find(l => l.id === st.lang);
      $('langBtn').innerHTML = FLAGS[cur.id] + '<span>' + cur.code + '</span>' +
        '<svg class="chev" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.4" stroke-linecap="round"><path d="m6 9 6 6 6-6"/></svg>';
      $('langMenu').setAttribute('aria-label', t('lang'));
      $('langMenu').innerHTML = LANGS.map(l =>
        '<li role="option" data-l="' + l.id + '" aria-selected="' + (l.id === st.lang) + '">' + FLAGS[l.id] +
        '<b>' + l.code + '</b><span class="nm">' + l.name + '</span>' +
        '<svg class="ck" width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.6" stroke-linecap="round"><path d="m5 12 5 5 9-10"/></svg></li>').join('');
    }

    function render() {
      renderStatic();
      $('view').setAttribute('aria-label', t('view'));
      $('view').innerHTML = [['flat', 'vFlat'], ['tree', 'vTree']].map(([v, k]) => '<button class="' + (st.view === v ? 'on' : '') + '" data-v="' + v + '">' + esc(t(k)) + '</button>').join('');
      $('groups').innerHTML = G.map(g => '<button class="' + (g === st.g ? 'on' : '') + '" data-g="' + g + '">' + g + '<span class="n">' +
        (flat() ? DATA.flat[g].items.length : DATA.groups[g].length) + '</span></button>').join('');
      document.querySelector('.machines').hidden = flat();
      const q0 = low(st.q.trim());
      if (flat()) { renderFlat(q0); return; }
      const list = DATA.groups[st.g], cur = machine();
      $('machines').innerHTML = list.map((m, i) => '<button class="' + (m === cur ? 'on' : '') + '" data-i="' + i + '">' + esc(m.m) + '<span class="n">' + m.items.length + '</span></button>').join('');
      $('title').textContent = cur.m;
      const q = low(st.q.trim());
      $('note').textContent = ''; $('extra').innerHTML = '';
      const rows = cur.items.map((it, i) => ({ it, i })).filter(({ it }) => matches(it, q) || subMatches(it.c, q));
      const komple = cur.items.filter(it => hasSub(it.c)).length;
      const foto = cur.items.filter(it => it.p).length;
      $('stats').innerHTML = '<span><b>' + cur.items.length + '</b>' + esc(t('items')) + '</span><span><b>' + komple + '</b>' + esc(t('komple')) +
        '</span><span><b>' + foto + '</b>' + esc(t('photos')) + '</span>' +
        (cur.x ? '<span title="' + esc(t('excluded')) + '"><b>' + cur.x + '</b>' + esc(t('hidden')) + '</span>' : '') +
        (q ? '<span><b>' + rows.length + '</b>' + esc(t('matches')) + '</span>' : '');
      if (!rows.length) { $('bom').innerHTML = '<div class="empty">' + esc(t('empty')) + '</div>'; return; }
      $('bom').innerHTML = '<table><thead><tr><th class="no">#</th><th></th><th>' + esc(t('hCode')) + '</th><th>' + esc(t('hName')) +
        '</th><th class="full">' + esc(t('hFull')) + '</th><th style="text-align:right">' + esc(t('hQty')) + '</th><th class="unit">' + esc(t('hUnit')) + '</th></tr></thead><tbody>' +
        rows.map(({ it, i }) => {
          const key = String(i), sub = hasSub(it.c), open = sub && (st.open.has(key) || (q && subMatches(it.c, q)));
          return '<tr><td class="no">' + (i + 1) + '</td><td class="ph">' + thumb(it.p, it.c, it.n) + '</td>' +
            '<td class="code mono">' + hl(it.c, q) + '</td>' +
            '<td class="name">' + (sub ? toggle(key, open) : '') + hl(it.n, q) + (sub ? subTag(it.c) : '') + '</td>' +
            fullCell(it.c, q) +
            '<td class="qty">' + esc(it.q) + '</td><td class="unit">' + esc(unit(it.u)) + '</td></tr>' +
            (open ? subRows(it.c, q, 1, key + '>' + it.c) : '');
        }).join('') + '</tbody></table>';
    }

    // Açık/kapalı anahtarı: kök satır sırası + '>' + kod zinciri
    $('view').onclick = e => { const b = e.target.closest('button'); if (!b) return; st.view = b.dataset.v; store.set('view', st.view); render(); };
    $('extra').addEventListener('click', e => { const im = e.target.closest('img.thumb'); if (im) showImg(im); });
    function showImg(im) {
      const f = full(im.dataset.c);
      $('lbImg').src = IMGS[im.dataset.p];
      $('lbCap').innerHTML = '<b class="mono">' + esc(im.dataset.c) + '</b>' + esc(im.dataset.n) + (f ? '<div class="capfull">' + esc(f) + '</div>' : '');
      $('lb').showModal();
    }
    $('bom').addEventListener('click', e => {
      const tg = e.target.closest('.tg');
      if (tg) {
        const k = tg.dataset.k;
        st.open.has(k) ? st.open.delete(k) : st.open.add(k);
        render(); return;
      }
      const im = e.target.closest('img.thumb');
      if (im) showImg(im);
    });
    $('lb').onclick = () => $('lb').close();
    $('groups').onclick = e => { const b = e.target.closest('button'); if (!b) return; st.g = b.dataset.g; st.open.clear(); store.set('g', st.g); render(); };
    $('machines').onclick = e => { const b = e.target.closest('button'); if (!b) return; st.m[st.g] = +b.dataset.i; st.open.clear(); store.set('m', JSON.stringify(st.m)); render(); };
    $('q').oninput = e => { st.q = e.target.value; render(); };
    $('printBtn').onclick = () => window.print();
    $('csvBtn').onclick = () => {
      const cur = machine(), cell = v => '"' + String(v ?? '').replace(/"/g, '""') + '"';
      if (flat()) {
        const F = DATA.flat[st.g];
        const rows = [[t('csvNo'), t('hCode'), t('hName'), t('hFull'), t('hQty'), t('hSrc')].map(cell).join(';')]
          .concat(F.items.map((it, i) => [i + 1, it.c, it.n, full(it.c), 1, it.s.join(' | ') || t('direct')].map(cell).join(';')));
        return download(rows, 'BOM ' + st.g + ' ' + t('vFlat') + ' ' + st.lang.toUpperCase() + '.csv');
      }
      const lines = [[t('csvNo'), t('hCode'), t('hName'), t('hFull'), t('hQty'), t('hUnit'), t('csvParent')].map(cell).join(';')];
      const addSubs = (c, parent, seen) => { if (!hasSub(c) || seen.has(c)) return; seen.add(c);
        DATA.subs[c].forEach(s => { lines.push(['', s.c, s.n, full(s.c), '', '', parent].map(cell).join(';')); addSubs(s.c, s.c, new Set(seen)); }); };
      cur.items.forEach((it, i) => { lines.push([i + 1, it.c, it.n, full(it.c), it.q, unit(it.u), ''].map(cell).join(';')); addSubs(it.c, it.c, new Set()); });
      download(lines, 'BOM ' + cur.m + ' ' + st.lang.toUpperCase() + '.csv');
    };
    function download(lines, name) {
      const blob = new Blob(['﻿' + lines.join('\r\n')], { type: 'text/csv;charset=utf-8' });
      const a = document.createElement('a'); a.href = URL.createObjectURL(blob); a.download = name; a.click();
      setTimeout(() => URL.revokeObjectURL(a.href), 1000);
    }

    // Dil menüsü
    function setLang(id) { st.lang = id; store.set('lang', id); render(); }
    function toggleLang(open) {
      const menu = $('langMenu');
      open = open ?? menu.hidden;
      menu.hidden = !open;
      $('langBtn').setAttribute('aria-expanded', String(open));
      if (open) { [...menu.children].forEach(li => li.classList.toggle('focus', li.dataset.l === st.lang)); menu.focus(); }
    }
    $('langBtn').onclick = e => { e.stopPropagation(); toggleLang(); };
    $('langMenu').onclick = e => { const li = e.target.closest('li'); if (!li) return; toggleLang(false); setLang(li.dataset.l); $('langBtn').focus(); };
    $('langMenu').onkeydown = e => {
      const items = [...$('langMenu').children];
      let i = items.findIndex(li => li.classList.contains('focus'));
      if (e.key === 'ArrowDown' || e.key === 'ArrowUp') {
        e.preventDefault(); i = (i + (e.key === 'ArrowDown' ? 1 : items.length - 1)) % items.length;
        items.forEach((li, j) => li.classList.toggle('focus', j === i));
      } else if (e.key === 'Enter' || e.key === ' ') {
        e.preventDefault(); if (i >= 0) { toggleLang(false); setLang(items[i].dataset.l); $('langBtn').focus(); }
      } else if (e.key === 'Escape' || e.key === 'Tab') { toggleLang(false); $('langBtn').focus(); }
    };
    document.addEventListener('click', e => { if (!$('lang').contains(e.target)) toggleLang(false); });
    document.addEventListener('keydown', e => {
      if (e.key === '/' && !/^(INPUT|SELECT|TEXTAREA)$/.test(document.activeElement.tagName)) { e.preventDefault(); $('q').focus(); }
    });
    const theme = store.get('theme');
    if (theme) document.documentElement.dataset.theme = theme;
    $('themeBtn').onclick = () => {
      const dark = document.documentElement.dataset.theme ? document.documentElement.dataset.theme === 'dark' : matchMedia('(prefers-color-scheme: dark)').matches;
      const next = dark ? 'light' : 'dark'; document.documentElement.dataset.theme = next; store.set('theme', next);
    };
    render();
  </script>
</body>
</html>
'''


def shared_lang_assets() -> tuple[str, str]:
    """Bayraklar ve dil menüsü stili kritik parça sayfasıyla aynı olsun diye oradan alınır."""
    tpl = bks.HTML_TEMPLATE
    flags = re.search(r"const FLAGS = \{.*?\n    \};", tpl, re.S).group(0)
    css = re.search(r"    /\* Dil seçimi \*/.*?(?=  </style>)", tpl, re.S).group(0)
    return flags, css


def main() -> None:
    data, imgs = build_data()
    flags, lang_css = shared_lang_assets()
    html = (HTML_TEMPLATE.replace("__DATA__", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
            .replace("__IMGS__", json.dumps(imgs))
            .replace("__FLAGS__", flags)
            .replace("__LANGCSS__", lang_css))
    with open(OUT_PATH, "w", encoding="utf-8") as f:
        f.write(html)
    summary = {g: [(m["m"], len(m["items"])) for m in data["groups"][g]] for g in GROUP_ORDER}
    print("BOM OK:", summary, "| foto:", len(imgs), "| tam ad:", len(data["full"]), "/", len(data["full"]) + len(data["missing"]))


if __name__ == "__main__":
    main()
